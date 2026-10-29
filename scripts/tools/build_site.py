#!/usr/bin/env python3
"""Baut die Vercel-Seite (site/) aus der Beispiel-Präsentation, der Layout-Galerie und den Download-Dateien.

Aufruf (nach make_template.py --gallery und build_deck.py, siehe README):
    python3 scripts/tools/build_site.py
Erzeugt site/index.html, site/slides/…, site/downloads/…, Impressum, Datenschutz und Favicon.
CSS, JS und Schriften liegen fest in site/assets/. Die Login-Seite steht in middleware.js.
"""
import hashlib
import html
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from PIL import Image

REPO = Path(__file__).resolve().parents[2]
SKILL = REPO / "skill" / "ecowerth-campus"
sys.path.insert(0, str(SKILL / "scripts"))

from pptx import Presentation  # noqa: E402
from pptx.enum.shapes import PP_PLACEHOLDER  # noqa: E402

from campus.layouts import LAYOUTS  # noqa: E402
from campus.render import render_pdf  # noqa: E402
from campus.tokens import PH_IDX  # noqa: E402

SITE = REPO / "site"
OUT = REPO / "out"
DECK_PLAN = SKILL / "examples" / "ecowerth-praesentation.json"
DECK = OUT / "ecowerth-praesentation.pptx"
GALLERY = SKILL / "assets" / "template" / "ecowerth-campus-template.pptx"
FULL_W, THUMB_W = 1600, 400
STAND = "29. Oktober 2026"
ASSET_VERSION = "1"  # Gleich dem ?v= der Schriften in site.css; erhöhen, wenn Browser alte Assets festhalten.


# ---------------------------------------------------------------- Inhalte aus dem Deck

def _placeholder_text(slide, idx: int | None = None, title: bool = False) -> str:
    for ph in slide.placeholders:
        is_title = ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE)
        if (title and is_title) or (not title and ph.placeholder_format.idx == idx):
            if ph.has_text_frame:
                return " ".join(p.text for p in ph.text_frame.paragraphs if p.text).strip()
    return ""


def _active_chapter(slide) -> int | None:
    """Aktives Kapitel aus der Kopfzeile (nav-N-active) oder dem Kapitelring (ring-seg-N-active)."""
    for s in slide.shapes:
        if s.name.endswith("-active") and s.name.startswith(("nav-", "ring-seg-")):
            return int(s.name.split("-")[-2])
    return None


def slide_records(pptx: Path) -> list[dict]:
    """Nummer, Titel, Kicker, aktives Kapitel und Sprechertext je Folie."""
    records = []
    for n, slide in enumerate(Presentation(str(pptx)).slides, 1):
        title = (_placeholder_text(slide, title=True) or _placeholder_text(slide, PH_IDX["Quote"])
                 or _placeholder_text(slide, PH_IDX["Question"]) or _placeholder_text(slide, PH_IDX["Statement"]))
        notes = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
        records.append({"n": n, "title": title, "kicker": _placeholder_text(slide, PH_IDX["Kicker"]),
                        "chapter": _active_chapter(slide), "notes": notes, "layout": slide.slide_layout.name})
    return records


# ---------------------------------------------------------------- Bilder

def render_images(pptx: Path, out_dir: Path) -> list[tuple[str, str]]:
    """Rendert alle Folien als JPEG (1600 px) und Vorschau (400 px). Gibt relative Pfade zurück."""
    (out_dir / "thumb").mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as tmp:
        pdf = render_pdf(pptx, Path(tmp))
        subprocess.run(["pdftoppm", "-png", "-scale-to-x", str(FULL_W), "-scale-to-y", "-1", str(pdf),
                        str(Path(tmp) / "p")], check=True)
        pages = sorted(Path(tmp).glob("p-*.png"))
        paths = []
        for i, page in enumerate(pages, 1):
            im = Image.open(page).convert("RGB")
            full, thumb = out_dir / f"{i:02d}.jpg", out_dir / "thumb" / f"{i:02d}.jpg"
            im.save(full, "JPEG", quality=84, optimize=True, progressive=True)
            im.resize((THUMB_W, round(im.height * THUMB_W / im.width)), Image.LANCZOS).save(
                thumb, "JPEG", quality=80, optimize=True)
            paths.append((full.relative_to(SITE).as_posix(), thumb.relative_to(SITE).as_posix()))
    return paths


def versioned(root: Path, rel: str) -> str:
    """Link mit Inhalts-Kennung: ein geändertes Bild bekommt eine neue Adresse, alte Caches greifen nicht."""
    return f"{rel}?v={hashlib.sha1((root / rel).read_bytes()).hexdigest()[:10]}"


# ---------------------------------------------------------------- Marke

def logo_svg(css_class: str = "logo") -> str:
    """ECOWERTH-Logo als SVG; Wortmarke in currentColor, Trennfuge (class="sep") per CSS in der Hintergrundfarbe."""
    svg = (SKILL / "assets" / "brand" / "logo.svg").read_text(encoding="utf-8")
    svg = svg[svg.index("<svg"):]
    svg = svg.replace(' width="427.71" height="71.21"', "")
    return svg.replace("<svg", f'<svg class="{css_class}" role="img" aria-label="ECOWERTH"', 1)


def logo_file_svg() -> str:
    """Logo als eigenständige Datei für die Login-Seite: Wortmarke Dunkelblau, Trennfuge in der Seitenfarbe."""
    svg = (SKILL / "assets" / "brand" / "logo.svg").read_text(encoding="utf-8")
    return svg.replace('fill="currentColor"', 'fill="#043751"').replace('class="sep" fill="#FFFFFF"', 'fill="#F4F6F8"')


def favicon_svg() -> str:
    """Das Signet allein (drei Balken, zwei Scherflächen) als Favicon."""
    svg = (SKILL / "assets" / "brand" / "logo.svg").read_text(encoding="utf-8")
    paths = [line for line in svg.splitlines() if line.startswith("<path") and "currentColor" not in line
             and 'class="sep"' not in line]
    return ('<svg xmlns="http://www.w3.org/2000/svg" viewBox="64 119 76 76">'
            '<rect x="64" y="119" width="76" height="76" rx="8" fill="#FFFFFF"/>' + "".join(paths) + "</svg>\n")


# ---------------------------------------------------------------- HTML

def _size(path: Path) -> str:
    return f"{path.stat().st_size / 1_000_000:.1f} MB".replace(".", ",")


HEAD = """<!doctype html>
<html lang="de">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<meta name="robots" content="noindex, nofollow">
<title>{title}</title>
<meta name="description" content="{description}">
<link rel="preload" href="/assets/fonts/archivo-600.woff2?v={v}" as="font" type="font/woff2" crossorigin>
<link rel="stylesheet" href="/assets/site.css?v={v}">
<link rel="icon" href="/favicon.svg" type="image/svg+xml">
</head>"""


def _footer() -> str:
    return f"""<footer class="site">
  <div class="wrap fgrid">
    {logo_svg()}
    <p>ECOWERTH Campus – Präsentationsvorlage für die ecowerth GmbH. Gestaltung und Umsetzung: Grafitecture. Stand: {STAND}.</p>
    <nav class="flegal" aria-label="Rechtliches"><a href="/impressum">Impressum</a><a href="/datenschutz">Datenschutz</a></nav>
  </div>
</footer>"""


def write_index(out_dir: Path, deck: dict, slides: list[dict], layouts: list[dict], downloads: list[dict]) -> Path:
    e = html.escape
    chapters = "".join(
        f'<li><button type="button" class="chapter" data-chapter="{i}" aria-label="Kapitel {i}: {e(name)}">'
        f'<span class="num">{i:02d}</span><span class="name">{e(name)}</span></button></li>'
        for i, name in enumerate(deck["chapters"], 1))
    thumbs = "".join(
        f'<li><button type="button" class="thumb" data-index="{i}" aria-label="Folie {s["n"]}: {e(s["title"])}">'
        f'<img src="/{s["thumb"]}" alt="" width="400" height="225" loading="lazy" decoding="async"></button></li>'
        for i, s in enumerate(slides))
    layout_items = "".join(
        f'<li><button type="button" class="layout" data-src="/{lay["src"]}" data-label="{e(lay["name"])}">'
        f'<img src="/{lay["thumb"]}" alt="Layout {e(lay["name"])}" width="400" height="225" loading="lazy" decoding="async">'
        f'<span class="layout-name">{e(lay["name"])}</span><span class="layout-key">{e(lay["key"])}</span></button></li>'
        for lay in layouts)
    files = "".join(
        f'<li><a class="file" href="/{d["href"]}" download><span class="file-label">{e(d["label"])}</span>'
        f'<span class="file-meta">{d["kind"]} · {d["size"]}</span><span class="file-action">Herunterladen</span></a></li>'
        for d in downloads)
    first = slides[0]
    data = json.dumps({"chapters": deck["chapters"], "slides": slides}, ensure_ascii=False).replace("</", "<\\/")
    count = len({lay["name"].split(" mit ")[0].split(" (")[0] for lay in layouts})
    page = HEAD.format(title=f"ECOWERTH Campus · {e(deck['title'])}", v=ASSET_VERSION,
                       description="ECOWERTH Campus: PowerPoint-Vorlage in der ECOWERTH-CI mit Beispiel-Präsentation.")
    page += f"""
<body>
<a class="skip" href="#viewer">Zur Präsentation springen</a>
<header class="site">
  <div class="wrap bar">
    <a class="brand" href="/" aria-label="ECOWERTH Campus, Start">{logo_svg()}<span class="campus">Campus</span></a>
    <nav class="main" aria-label="Bereiche"><a href="#folien">Folien</a><a href="#layouts">Layouts</a><a href="#dateien">Dateien</a></nav>
    <a class="logout" href="/abmelden">Abmelden</a>
  </div>
</header>
<main>
  <section class="stage on-dark" id="folien" aria-labelledby="titel">
    <div class="wrap">
      <div class="intro">
        <p class="eyebrow">ECOWERTH Campus · Präsentationsvorlage</p>
        <h1 id="titel">{e(deck["title"])}</h1>
        <p>Die PowerPoint-Vorlage im Erscheinungsbild der ECOWERTH-Website: {count} Layouts und eine Beispiel-Präsentation mit den Inhalten der Website. Blättern mit den Pfeiltasten oder über die Vorschau.</p>
      </div>
      <ol class="chapters" aria-label="Kapitel">{chapters}</ol>
      <figure class="viewer" id="viewer" tabindex="-1">
        <div class="screen"><img id="slide" src="/{first["src"]}" alt="Folie 1: {e(first["title"])}" width="1600" height="900" fetchpriority="high"></div>
        <figcaption class="controls">
          <button type="button" class="step" id="prev" aria-label="Vorherige Folie"><svg viewBox="0 0 24 24" aria-hidden="true"><path d="M15 5l-7 7 7 7"/></svg></button>
          <span class="counter" aria-live="polite"><span id="current">01</span> / {len(slides):02d}</span>
          <button type="button" class="step" id="next" aria-label="Nächste Folie"><svg viewBox="0 0 24 24" aria-hidden="true"><path d="M9 5l7 7-7 7"/></svg></button>
          <span class="where" id="where"></span>
          <button type="button" class="notes-toggle" id="notes-toggle" aria-expanded="false" aria-controls="notes">Sprechertext anzeigen</button>
        </figcaption>
        <div class="notes" id="notes" hidden><p id="notes-text"></p></div>
      </figure>
      <ol class="thumbs" aria-label="Alle Folien">{thumbs}</ol>
    </div>
  </section>
  <section class="paper" id="layouts" aria-labelledby="layouts-titel">
    <div class="wrap">
      <div class="sec-head">
        <p class="eyebrow">Vorlage</p>
        <h2 id="layouts-titel">Die {count} Layouts</h2>
        <p>Aus diesen Layouts entsteht jede Folie. Die Beispiele zeigen Inhalte der ECOWERTH-Website. Der Schlüssel unter dem Namen hilft bei Korrekturen, etwa „Folie 7 als compare“.</p>
      </div>
      <ul class="layouts">{layout_items}</ul>
    </div>
  </section>
  <section class="paper soft" id="dateien" aria-labelledby="dateien-titel">
    <div class="wrap">
      <div class="sec-head">
        <p class="eyebrow">Downloads</p>
        <h2 id="dateien-titel">Dateien</h2>
      </div>
      <ul class="files">{files}</ul>
      <p class="hint"><b>Zum Bearbeiten in PowerPoint</b> die Schriften der Website installieren: <a href="https://fonts.google.com/specimen/Archivo" rel="noreferrer" target="_blank">Archivo</a>, <a href="https://fonts.google.com/specimen/IBM+Plex+Sans" rel="noreferrer" target="_blank">IBM Plex Sans</a> und <a href="https://fonts.google.com/specimen/IBM+Plex+Mono" rel="noreferrer" target="_blank">IBM Plex Mono</a> (kostenlos, SIL Open Font License). Die PDF-Dateien enthalten die Schriften bereits.</p>
    </div>
  </section>
</main>
{_footer()}
<dialog class="lightbox" id="lightbox" aria-labelledby="lightbox-label">
  <form method="dialog"><button class="close" aria-label="Schließen">×</button></form>
  <img id="lightbox-img" src="" alt="">
  <p id="lightbox-label"></p>
</dialog>
<script type="application/json" id="deck-data">{data}</script>
<script src="/assets/site.js?v={ASSET_VERSION}" defer></script>
</body>
</html>
"""
    out_dir.mkdir(parents=True, exist_ok=True)
    path = out_dir / "index.html"
    path.write_text(page, encoding="utf-8")
    return path


# ---------------------------------------------------------------- Impressum und Datenschutz



TODO = '<span class="todo">PLATZHALTER — {}</span>'
CONTACT = ('Telefon: <a href="tel:+4915110855833">+49 151 10 855 833</a><br>'
           'E-Mail: <a href="mailto:florian.werthmann@ecowerth.de">florian.werthmann@ecowerth.de</a>')

IMPRESSUM = f"""<p class="eyebrow">Rechtliches</p>
<h1>Impressum</h1>
<div class="legal">
<h3>Angaben gemäß § 5 DDG</h3>
<p>ecowerth GmbH<br>Goldleite 9<br>97234 Reichenberg<br>Deutschland</p>
<h3>Vertreten durch</h3>
<p>Florian Werthmann, Geschäftsführer</p>
<h3>Kontakt</h3>
<p>{CONTACT}</p>
<h3>Registereintrag</h3>
<p>Eintragung im Handelsregister<br>Registergericht und Registernummer: {TODO.format("Amtsgericht und HRB-Nummer der ecowerth GmbH")}</p>
<h3>Umsatzsteuer-ID</h3>
<p>Umsatzsteuer-Identifikationsnummer gemäß § 27a Umsatzsteuergesetz: {TODO.format("USt-IdNr. der ecowerth GmbH")}</p>
<h3>Aufsichtsbehörde</h3>
<p>Zuständige Aufsichtsbehörde: {TODO.format("nur falls eine Tätigkeit erlaubnispflichtig ist (z. B. § 54 KrWG oder Zulassung als Organisation für Herstellerverantwortung nach BattDG): zuständige Behörde laut Bescheid; sonst entfällt dieser Absatz")}</p>
<h3>Verantwortlich für den Inhalt nach § 18 Abs. 2 MStV</h3>
<p>Florian Werthmann<br>Goldleite 9<br>97234 Reichenberg</p>
<h3>Verbraucherstreitbeilegung</h3>
<p>Wir sind nicht bereit und nicht verpflichtet, an Streitbeilegungsverfahren vor einer Verbraucherschlichtungsstelle teilzunehmen.</p>
<h3>Hinweis zu den Inhalten</h3>
<p>Diese Seite zeigt die Präsentationsvorlage ECOWERTH Campus. Wir erstellen die Inhalte mit großer Sorgfalt, übernehmen für ihre Richtigkeit, Vollständigkeit und Aktualität jedoch keine Gewähr. Die Informationen zu gesetzlichen Pflichten dienen der allgemeinen Orientierung und ersetzen keine Beratung im Einzelfall.</p>
<h3>Urheberrecht</h3>
<p>Die Inhalte dieser Seite sind, soweit schutzfähig, urheberrechtlich geschützt. Eine Verwendung über die gesetzlich erlaubten Grenzen hinaus bedarf der vorherigen Zustimmung des jeweiligen Rechteinhabers.</p>
<h3>Datenschutz</h3>
<p>Wie wir mit personenbezogenen Daten umgehen, erläutert unsere <a href="/datenschutz">Datenschutzerklärung</a>.</p>
<p class="stand">Stand: Oktober 2026</p>
</div>"""

DATENSCHUTZ = f"""<p class="eyebrow">Rechtliches</p>
<h1>Datenschutzerklärung</h1>
<div class="legal">
<p>Kurz gesagt: Diese Seite setzt nach Ihrer Anmeldung ein einziges, technisch notwendiges Cookie, verfolgt Sie nicht und lädt nichts von Servern Dritter nach. Sie wird beim US-Anbieter Vercel gehostet. Was dabei mit Ihren Daten geschieht, erklären wir hier.</p>
<h3>Verantwortlicher</h3>
<p>Verantwortlich für die Verarbeitung personenbezogener Daten auf dieser Seite ist:</p>
<p>ecowerth GmbH<br>Goldleite 9<br>97234 Reichenberg<br>Deutschland<br>vertreten durch den Geschäftsführer Florian Werthmann</p>
<p>{CONTACT}</p>
<p>Datenschutzbeauftragter: {TODO.format("Kontaktdaten, sofern ein Datenschutzbeauftragter benannt ist; sonst entfällt dieser Absatz")}</p>
<h3>Überblick</h3>
<p>Diese Seite verwendet keine Analyse- oder Tracking-Werkzeuge und keine Social-Media-Plugins, auch nicht Vercel Web Analytics oder Vercel Speed Insights. Schriften, Bilder und alle übrigen Bestandteile werden mit der Seite selbst ausgeliefert; beim Aufruf werden keine Inhalte von Servern Dritter geladen. Links zu Google Fonts im Bereich „Dateien“ öffnen erst, wenn Sie sie anklicken.</p>
<h3>Hosting bei Vercel</h3>
<p>Diese Seite wird bei der Vercel Inc., 440 N Barranca Avenue #4133, Covina, CA 91723, USA („Vercel“) gehostet. Vercel liefert die Seite über ein weltweites Netz von Rechenzentren aus; Anfragen werden in der Regel von einem Standort in Ihrer Nähe beantwortet.</p>
<p>Bei jedem Aufruf verarbeitet Vercel die Verbindungsdaten, die Ihr Browser automatisch übermittelt:</p>
<ul>
<li>IP-Adresse und eine daraus abgeleitete ungefähre Standortangabe</li>
<li>Datum und Uhrzeit des Zugriffs</li>
<li>aufgerufene Adresse, Statuscode und übertragene Datenmenge</li>
<li>die zuvor besuchte Seite (Referrer), sofern Ihr Browser sie übermittelt</li>
<li>Browsertyp, Betriebssystem und weitere Angaben zur Systemkonfiguration</li>
</ul>
<p>Wir benötigen diese Daten, um die Seite auszuliefern, ihre Stabilität und Sicherheit zu gewährleisten, etwa zur Abwehr von Angriffen, und um Fehler zu analysieren. Rechtsgrundlage ist Art. 6 Abs. 1 lit. f DSGVO; unser berechtigtes Interesse liegt in einer sicheren und zuverlässigen Bereitstellung der Seite. Wir werten die Daten nicht zu anderen Zwecken aus und führen sie nicht mit anderen Daten zusammen. In unserem Vercel-Konto sind die Zugriffsprotokolle höchstens einen Tag lang abrufbar. Vercel selbst bewahrt Verbindungsdaten nach eigenen Angaben nur so lange auf, wie es für Betrieb und Sicherheit des Dienstes erforderlich ist.</p>
<p>Vercel verarbeitet die Daten in unserem Auftrag auf Grundlage eines Vertrags zur Auftragsverarbeitung (Data Processing Addendum) nach Art. 28 DSGVO. {TODO.format("gilt erst, wenn das Vercel-Konto im Pro-Tarif läuft (den Vertrag bietet Vercel nur für Pro und Enterprise an). Bleibt das Projekt im Vercel-Team von Grafitecture, zusätzlich einen AV-Vertrag zwischen ecowerth und Grafitecture schließen und diesen Absatz anpassen")}</p>
<p>Vercel hat seinen Sitz in den USA; die Daten können auch dort verarbeitet werden. Für US-Unternehmen, die nach dem EU-U.S. Data Privacy Framework zertifiziert sind, hat die Europäische Kommission ein angemessenes Schutzniveau festgestellt (Durchführungsbeschluss (EU) 2023/1795). Vercel ist danach zertifiziert; die Übermittlung stützt sich daher auf Art. 45 DSGVO. Zusätzlich sind im Auftragsverarbeitungsvertrag die Standardvertragsklauseln der Europäischen Kommission vereinbart (Art. 46 Abs. 2 lit. c DSGVO).</p>
<p>Weitere Informationen: <a href="https://vercel.com/legal/privacy-notice" rel="noreferrer" target="_blank">Datenschutzhinweise von Vercel</a> · <a href="https://vercel.com/legal/dpa" rel="noreferrer" target="_blank">Auftragsverarbeitungsvertrag von Vercel</a></p>
<h3>Cookie für den Zugang</h3>
<p>Die Präsentationsseite ist passwortgeschützt. Nach der Anmeldung setzt sie das Cookie <b>ew_campus</b>. Es enthält einen aus dem Passwort abgeleiteten Prüfwert, damit Sie das Passwort nicht bei jedem Aufruf neu eingeben müssen, und wird nur an diese Seite übertragen. Setzen Sie den Haken bei „Angemeldet bleiben“, wird es 30 Tage gespeichert, sonst nur bis zum Schließen des Browsers. Über „Abmelden“ wird es sofort gelöscht.</p>
<p>Das Speichern und Auslesen dieses Cookies ist unbedingt erforderlich, um den von Ihnen gewünschten Zugang bereitzustellen; eine Einwilligung ist dafür nicht nötig (§ 25 Abs. 2 Nr. 2 TDDDG). Rechtsgrundlage für die damit verbundene Verarbeitung ist Art. 6 Abs. 1 lit. f DSGVO. Sie können das Cookie jederzeit in den Einstellungen Ihres Browsers löschen.</p>
<p>Nur wenn der Angriffsschutz von Vercel eine Anfrage für verdächtig hält, kann er eine automatische Browserprüfung vorschalten. Nach bestandener Prüfung legt er für eine Stunde eine Sitzung in Ihrem Browser an. Das ist zur Abwehr von Angriffen technisch erforderlich (§ 25 Abs. 2 Nr. 2 TDDDG, Art. 6 Abs. 1 lit. f DSGVO).</p>
<h3>Kontakt per E-Mail oder Telefon</h3>
<p>Wenn Sie uns per E-Mail oder Telefon kontaktieren, verarbeiten wir die Angaben, die Sie uns dabei machen, etwa Name, Unternehmen, Kontaktdaten und Ihr Anliegen, um Ihre Anfrage zu bearbeiten. Rechtsgrundlage ist Art. 6 Abs. 1 lit. b DSGVO, soweit Sie selbst Vertragspartner sind oder werden wollen. Schreiben Sie uns für ein Unternehmen oder in anderen Fällen, ist Rechtsgrundlage Art. 6 Abs. 1 lit. f DSGVO; unser berechtigtes Interesse liegt in der Beantwortung von Anfragen und der Pflege von Geschäftskontakten.</p>
<p>Wir löschen diese Daten, sobald sie für die Bearbeitung nicht mehr erforderlich sind. Gesetzliche Aufbewahrungspflichten, insbesondere nach Handels- und Steuerrecht (§ 257 HGB, § 147 AO), bleiben unberührt.</p>
<h3>Empfänger</h3>
<p>Personenbezogene Daten geben wir nur weiter, soweit es für die genannten Zwecke erforderlich ist: an unseren Hosting-Anbieter Vercel sowie an IT-Dienstleister, die diese Seite und unsere E-Mail-Postfächer technisch betreuen. Sie handeln in unserem Auftrag und sind an unsere Weisungen gebunden. Wir verkaufen keine Daten und geben sie nicht zu Werbezwecken weiter.</p>
<h3>Pflicht zur Bereitstellung</h3>
<p>Sie sind weder gesetzlich noch vertraglich verpflichtet, uns personenbezogene Daten bereitzustellen. Ohne die Verbindungsdaten lässt sich die Seite jedoch nicht anzeigen, und ohne Kontaktdaten können wir Ihre Anfrage nicht beantworten. Eine automatisierte Entscheidungsfindung einschließlich Profiling (Art. 22 DSGVO) findet nicht statt.</p>
<h3>Ihre Rechte</h3>
<p>Sie haben uns gegenüber das Recht auf Auskunft (Art. 15 DSGVO), Berichtigung (Art. 16), Löschung (Art. 17), Einschränkung der Verarbeitung (Art. 18) und Datenübertragbarkeit (Art. 20). Eine erteilte Einwilligung können Sie jederzeit mit Wirkung für die Zukunft widerrufen (Art. 7 Abs. 3 DSGVO). Eine formlose Nachricht an die oben genannten Kontaktdaten genügt.</p>
<p class="legal-box"><b>Widerspruchsrecht nach Art. 21 DSGVO:</b> Verarbeiten wir Ihre Daten auf Grundlage von Art. 6 Abs. 1 lit. f DSGVO, können Sie dieser Verarbeitung jederzeit aus Gründen widersprechen, die sich aus Ihrer besonderen Situation ergeben. Wir verarbeiten die Daten dann nicht mehr, es sei denn, wir können zwingende schutzwürdige Gründe nachweisen, die Ihre Interessen, Rechte und Freiheiten überwiegen, oder die Verarbeitung dient der Geltendmachung, Ausübung oder Verteidigung von Rechtsansprüchen.</p>
<h3>Beschwerderecht</h3>
<p>Sie können sich bei einer Datenschutz-Aufsichtsbehörde beschweren, insbesondere in dem Mitgliedstaat Ihres Aufenthaltsorts, Ihres Arbeitsplatzes oder des mutmaßlichen Verstoßes (Art. 77 DSGVO). Für uns zuständig ist das Bayerische Landesamt für Datenschutzaufsicht, Promenade 18, 91522 Ansbach, <a href="https://www.lda.bayern.de" rel="noreferrer" target="_blank">www.lda.bayern.de</a>.</p>
<h3>Verschlüsselung</h3>
<p>Die Verbindung zu dieser Seite ist per TLS verschlüsselt. Sie erkennen das an „https://“ in der Adresszeile Ihres Browsers.</p>
<h3>Änderungen</h3>
<p>Wir passen diese Datenschutzerklärung an, sobald sich die Seite oder die Rechtslage ändert. Es gilt die jeweils hier veröffentlichte Fassung.</p>
<p class="stand">Stand: Oktober 2026</p>
</div>"""


def write_legal(out_dir: Path) -> list[Path]:
    """Impressum und Datenschutz: ohne Login erreichbar (siehe middleware.js), Gestaltung wie die Seite."""
    paths = []
    for slug, title, body in (("impressum", "Impressum", IMPRESSUM), ("datenschutz", "Datenschutzerklärung", DATENSCHUTZ)):
        page = HEAD.format(title=f"{title} · ECOWERTH Campus", v=ASSET_VERSION,
                           description=f"{title} der Präsentationsseite ECOWERTH Campus.")
        page += f"""
<body>
<header class="site">
  <div class="wrap bar">
    <a class="brand" href="/" aria-label="ECOWERTH Campus, Start">{logo_svg()}<span class="campus">Campus</span></a>
  </div>
</header>
<main class="legal-page"><div class="wrap">
{body}
</div></main>
{_footer()}
</body>
</html>
"""
        path = out_dir / f"{slug}.html"
        path.write_text(page, encoding="utf-8")
        paths.append(path)
    return paths


# ---------------------------------------------------------------- Ablauf

def _layout_records(gallery: Path, paths: list[tuple[str, str]]) -> list[dict]:
    """Je Galeriefolie: Layoutname, Schlüssel im Folienplan, Bild. Zweite Folien eines Layouts bekommen einen Zusatz."""
    key_by_name = {ldef.name: ldef.key for ldef in LAYOUTS}
    seen, layouts = set(), []
    for rec, (src, thumb) in zip(slide_records(gallery), paths):
        name = rec["layout"]
        if name in seen:
            name += " mit Fotos" if "foto" in rec["kicker"].lower() else " (Auflösung)"
        seen.add(rec["layout"])
        layouts.append({"name": name, "key": key_by_name.get(rec["layout"], ""), "src": src, "thumb": thumb})
    return layouts


def download_sources() -> list[tuple[str, Path, str, str]]:
    """Je Download: Beschriftung, Quelldatei, Dateiname auf der Seite, Dateiart."""
    return [
        ("Beispiel-Präsentation als PDF", OUT / "ecowerth-praesentation.pdf", "ecowerth-campus-zukunft-im-kreislauf.pdf",
         "PDF"),
        ("Beispiel-Präsentation als PowerPoint", DECK, "ecowerth-campus-zukunft-im-kreislauf.pptx", "PowerPoint"),
        ("Vorlage mit allen Layouts", GALLERY, "ecowerth-campus-vorlage.pptx", "PowerPoint"),
        ("Layout-Übersicht als PDF", OUT / "ecowerth-campus-layouts.pdf", "ecowerth-campus-layouts.pdf", "PDF"),
    ]


def main() -> None:
    plan = json.loads(DECK_PLAN.read_text(encoding="utf-8"))
    pres = plan["presentation"]
    for d in ("slides", "downloads"):
        shutil.rmtree(SITE / d, ignore_errors=True)
    render_pdf(DECK, OUT)                                          # out/ecowerth-praesentation.pdf
    with tempfile.TemporaryDirectory() as tmp:
        shutil.copyfile(render_pdf(GALLERY, Path(tmp)), OUT / "ecowerth-campus-layouts.pdf")
    records = slide_records(DECK)
    for rec, (src, thumb) in zip(records, render_images(DECK, SITE / "slides" / "praesentation")):
        rec.update(src=versioned(SITE, src), thumb=versioned(SITE, thumb))
        rec.pop("layout")
    gallery_paths = [(versioned(SITE, src), versioned(SITE, thumb))
                     for src, thumb in render_images(GALLERY, SITE / "slides" / "layouts")]
    layouts = _layout_records(GALLERY, gallery_paths)
    dl = SITE / "downloads"
    dl.mkdir(parents=True)
    downloads = []
    for label, src, name, kind in download_sources():
        shutil.copyfile(src, dl / name)
        downloads.append({"label": label, "href": f"downloads/{name}", "size": _size(dl / name), "kind": kind})
    (SITE / "favicon.svg").write_text(favicon_svg(), encoding="utf-8")
    public = SITE / "assets" / "public"
    public.mkdir(parents=True, exist_ok=True)
    (public / "logo.svg").write_text(logo_file_svg(), encoding="utf-8")
    (SITE / "robots.txt").write_text("User-agent: *\nDisallow: /\n", encoding="utf-8")
    write_legal(SITE)
    deck = {"title": pres["title"], "chapters": pres.get("chapters", [])}
    print(write_index(SITE, deck, records, layouts, downloads))


if __name__ == "__main__":
    main()
