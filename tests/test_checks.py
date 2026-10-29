import copy

import pytest

from campus.build import build_deck
from campus.checks import check_deck
from conftest import SKILL

BASE = {"presentation": {"title": "Prüfung", "chapters": ["Eins", "Zwei"]}, "slides": [
    {"type": "points", "chapter": 1, "title": "Kurz",
     "points": [{"lead": "Eins", "text": "Ein Satz."}, {"lead": "Zwei", "text": "Noch einer."}]}]}


@pytest.fixture
def plan():
    return copy.deepcopy(BASE)


def _issues(plan, template_path, tmp_path, m):
    out = tmp_path / "deck.pptx"
    build_deck(plan, template_path, out, tmp_path, m)
    return check_deck(out, m)


def test_clean_minimal(plan, template_path, tmp_path, m):
    assert _issues(plan, template_path, tmp_path, m) == []


def test_overflow(plan, template_path, tmp_path, m):
    plan["slides"][0]["points"][0]["text"] = "Viel zu langer Text, der über mehrere Zeilen läuft. " * 8
    plan["slides"][0]["points"] *= 2
    issues = _issues(plan, template_path, tmp_path, m)
    assert [(i.kind, i.shape) for i in issues] == [("overflow", "Points")]


def test_title_must_fit_one_line(plan, template_path, tmp_path, m):
    plan["slides"][0]["title"] = "Ein Titel, der deutlich zu lang für eine einzige Zeile auf der Folie ist"
    assert [(i.kind, i.shape) for i in _issues(plan, template_path, tmp_path, m)] == [("overflow", "Title")]


def test_glyph(plan, template_path, tmp_path, m):
    plan["slides"][0]["title"] = "Pfeil → Haken ✓"
    issues = _issues(plan, template_path, tmp_path, m)
    glyphs = [i for i in issues if i.kind == "glyph"]
    assert glyphs and "→" in glyphs[0].detail and "✓" in glyphs[0].detail


def test_font(plan, template_path, tmp_path, m):
    from pptx import Presentation

    out = tmp_path / "deck.pptx"
    build_deck(plan, template_path, out, tmp_path, m)
    prs = Presentation(out)
    title = next(ph for ph in prs.slides[0].placeholders if ph.placeholder_format.idx == 0)
    title.text_frame.paragraphs[0].runs[0].font.name = "Comic Sans MS"
    prs.save(out)
    assert any(i.kind == "font" and "Comic Sans MS" in i.detail for i in check_deck(out, m))


def test_long_word(plan, template_path, tmp_path, m):
    plan["slides"][0]["title"] = "Kreislaufwirtschaftsgesetzesdurchführungsverordnungsentwurfsbegründungsanlage"
    assert any(i.kind == "word" and i.shape == "Title" for i in _issues(plan, template_path, tmp_path, m))


def test_bounds(plan, template_path, tmp_path, m):
    from pptx import Presentation

    out = tmp_path / "deck.pptx"
    build_deck(plan, template_path, out, tmp_path, m)
    prs = Presentation(out)
    shape = prs.slides[0].shapes[-1]
    shape.left = prs.slide_width - 10
    prs.save(out)
    assert any(i.kind == "bounds" for i in check_deck(out, m))


@pytest.mark.parametrize("name", ["layout-gallery.json", "ecowerth-praesentation.json"])
def test_examples_are_clean(name, template_path, tmp_path, m):
    from campus.plan import load_plan

    out = tmp_path / "beispiel.pptx"
    report = build_deck(load_plan(SKILL / "examples" / name), template_path, out, SKILL / "examples", m)
    assert report.warnings == []
    assert check_deck(out, m) == []
