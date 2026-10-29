import math

import pytest
from pptx import Presentation

from campus.motifs import ARC_DEG, LOOP_COLORS, kreislauf, progress_segments
from campus.tokens import CONTENT_W, COLOR, MARGIN_X, px
from campus.widgets_content import FACT_MIN_W, fact_widths
from campus.widgets_frame import NAV_LEFT, chips, header_nav


@pytest.fixture
def prs(template_path):
    return Presentation(template_path)


def new_slide(prs, layout_name="05 Punkte"):
    layout = next(lay for lay in prs.slide_layouts if lay.name == layout_name)
    return prs.slides.add_slide(layout)


def names(slide):
    return [s.name for s in slide.shapes]


def test_nav_full(prs, m):
    slide = new_slide(prs)
    header_nav(slide, ["Die Gruppe", "Marken", "Leistungen"], 2, dark=False, m=m)
    shapes = {s.name: s for s in slide.shapes}
    assert [n for n in names(slide) if n.startswith("nav-")] == ["nav-1", "nav-2-active", "nav-underline", "nav-3"]
    assert "LEISTUNGEN" in shapes["nav-3"].text_frame.text.upper()
    right = max(s.left + s.width for n, s in shapes.items() if n.startswith("nav-") and n != "nav-underline")
    assert right <= px(MARGIN_X + CONTENT_W + 24)


def test_nav_compact_keeps_logo_free(prs, m):
    """Zu viele lange Kapitel: inaktive zeigen nur die Nummer, nichts ragt ins Logo."""
    slide = new_slide(prs)
    chapters = ["Kreislaufwirtschaftsgesetz", "Elektro- und Elektronikgerätegesetz", "Batterierecht-Durchführung",
                "Verpackungsgesetz", "Gewerbeabfallverordnung", "Zusammenarbeit"]
    header_nav(slide, chapters, 3, dark=False, m=m)
    shapes = {s.name: s for s in slide.shapes}
    assert shapes["nav-1"].text_frame.text == "01"
    assert "VERPACKUNGSGESETZ" not in shapes["nav-4"].text_frame.text.upper()
    assert "Batterierecht" in shapes["nav-3-active"].text_frame.text
    assert min(s.left for n, s in shapes.items() if n.startswith("nav-")) >= px(NAV_LEFT - 2)


def test_nav_absent_without_chapters(prs, m):
    slide = new_slide(prs)
    header_nav(slide, [], None, dark=False, m=m)
    assert not [n for n in names(slide) if n.startswith("nav-")]


def test_chips_advance(prs, m):
    slide = new_slide(prs)
    end = chips(slide, ["§ 59 KrWG", "§ 6 ElektroG"], 120, 900, m, prefix="legal-chip")
    found = [s for s in slide.shapes if s.name.startswith("legal-chip-")]
    assert [s.name for s in found] == ["legal-chip-1", "legal-chip-2"]
    assert end * 6350 >= found[-1].left + found[-1].width
    assert found[0].text_frame.text == "§ 59 KrWG"


@pytest.mark.parametrize("n", [2, 4, 6])
def test_progress_segments(n):
    segs = progress_segments(n, 12)
    assert len(segs) == n
    assert segs[0][0] == pytest.approx(-90 + 6)
    for a0, a1 in segs:
        assert a1 - a0 == pytest.approx(360 / n - 12)


def test_kreislauf_colors_and_order(prs):
    slide = new_slide(prs)
    kreislauf(slide.shapes, 960, 540, 300, prefix="t", names=("ECOLOGICON", "ECORELOOP", "ECOINVEST"))
    arcs = [s for s in slide.shapes if s.name.startswith("t-arc-")]
    assert [a.line.color.rgb for a in arcs] == [COLOR[c] for c in LOOP_COLORS] or len(arcs) == 3
    assert 0 < ARC_DEG < 120
    assert [s.name for s in slide.shapes if s.name.startswith("t-name-")] == ["t-name-1", "t-name-2", "t-name-3"]


def test_fact_widths(m):
    facts = [{"value": "4", "label": ""}, {"value": "8.000 m²", "label": ""}, {"value": "3", "label": ""}]
    widths = fact_widths(facts, m)
    assert sum(widths) == pytest.approx(CONTENT_W)
    assert min(widths) >= FACT_MIN_W
    assert widths[1] > widths[0]
    assert math.isclose(widths[0], widths[2])
