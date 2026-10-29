import copy
import json
import subprocess
import sys

import pytest
from pptx import Presentation

from campus.build import build_deck
from campus.checks import check_deck
from conftest import REPO, SKILL

PRES = {"title": "Zukunft im Kreislauf", "subtitle": "Kreislaufwirtschaft aus einer Hand.",
        "presenter": {"name": "Florian Werthmann", "role": "Geschäftsführer"}, "date": "Oktober 2026",
        "chapters": ["Die Gruppe", "Unsere Marken", "Leistungen", "Rechtsrahmen"]}
POINTS = [{"lead": "Beratung", "text": "Pflichtenanalyse und Umsetzung."},
          {"lead": "Systembetrieb", "text": "Rücknahme, Lagerung und Logistik."}]
ALL_TYPES = {"presentation": PRES, "slides": [
    {"type": "cover"},
    {"type": "agenda"},
    {"type": "chapter", "chapter": 2, "line": "Ein Dach, drei Marken."},
    {"type": "hook", "quote": "Welche Pflichten treffen uns?", "truth": {"text": "Es hängt an der Rolle."},
     "answer": {"text": "Wir ordnen zu und setzen um."}},
    {"type": "points", "chapter": 3, "title": "Leistungen", "points": POINTS, "legal": ["§ 59 KrWG"],
     "notes": "Sprechertext zu Folie fünf."},
    {"type": "points_photo", "chapter": 3, "title": "Umweltabteilung", "points": POINTS,
     "photo": {"brief": "Foto: Mitarbeiterin mit Tablet"}, "caption": "Im Betrieb"},
    {"type": "points_chart", "chapter": 2, "title": "ECO Park", "points": POINTS,
     "chart": {"title": "Flächen", "bars": [{"label": "Grundstück", "value": 100, "display": "8.000 m²",
                                             "highlight": True}, {"label": "Halle", "value": 22.5}]}},
    {"type": "statement", "big": "4", "statement": "Unternehmen unter einem Dach"},
    {"type": "facts", "title": "In Zahlen", "facts": [{"value": "4", "label": "Unternehmen"},
                                                       {"value": "8.000 m²", "label": "Standort"}]},
    {"type": "brands", "title": "Marken", "columns": [{"brand": "ecologicon", "items": ["Beratung"]},
                                                     {"brand": "ecoreloop"}, {"brand": "ecoinvest"}]},
    {"type": "checklist", "chapter": 3, "title": "Portfolio", "items": [f"Leistungsfeld {i}" for i in range(1, 9)]},
    {"type": "steps", "title": "Ablauf", "steps": [{"title": f"Schritt {i}", "text": "Tun."} for i in range(4)]},
    {"type": "compare", "title": "Vorher, nachher", "style": "neutral",
     "left": {"label": "Ändert sich", "items": ["Firmierung"]}, "right": {"label": "Bleibt", "items": ["Teams"]}},
    {"type": "matrix", "title": "Pflichten", "corner": "Akteur", "columns": ["KrWG", "ElektroG"],
     "rows": [{"label": "Importeure", "cells": [True, False]}, {"label": "Entsorger", "cells": [True, "teils"]}]},
    {"type": "scale", "title": "Zielgruppen", "tags": {"items": ["Hersteller", "Importeure"],
                                                       "highlight": ["Hersteller"]},
     "points": [{"lead": "Einstieg", "text": "Über die Matrix."}]},
    {"type": "photo", "title": "Kreisläufe schließen", "photo": {"brief": "Foto: Sortierhalle"}, "caption": "Halle"},
    {"type": "photo_pair", "title": "Zwei Boxen", "left": {"label": "Green Box", "photo": {"brief": "Foto: Box"}},
     "right": {"label": "Electro-Box", "photo": {"brief": "Foto: Box"}}},
    {"type": "photo_annotated", "title": "Electro-Box", "photo": {"brief": "Foto: Box"},
     "markers": [{"x": 0.3, "y": 0.4, "label": "Innenbehälter"}]},
    {"type": "case", "title": "Praxisfall", "story": "Ein Händler importiert Elektrogeräte.",
     "findings": [{"finding": "Importiert", "cause": "ElektroG § 6"}], "columns": ["Situation", "Pflicht"]},
    {"type": "quote", "quote": "Haltung gegenüber Ressourcen.", "name": "Florian Werthmann", "role": "Geschäftsführer"},
    {"type": "quiz", "question": "Welche Regelwerke?", "options": ["KrWG", "Alle drei", "Keine"], "answer": 1,
     "explanation": "Laut Matrix alle drei."},
    {"type": "summary", "title": "Aus einer Hand", "items": [{"title": "Ganzheitlich", "text": "Alles greift."}]},
    {"type": "closing", "title": "Lernen Sie uns kennen", "steps": [{"title": "Gespräch", "text": "Rolle schildern."}],
     "next": {"title": "Florian Werthmann", "text": "florian.werthmann@ecowerth.de"}},
]}
EXPECTED_LAYOUTS = ["01 Titel", "02 Agenda", "03 Kapitel", "04 Einstieg", "05 Punkte", "06 Punkte + Foto",
                    "07 Punkte + Grafik", "08 Aussage", "09 Kennzahlen", "10 Marken", "11 Checkliste", "12 Schritte",
                    "13 Vergleich", "14 Tabelle", "15 Skala", "16 Foto", "17 Zwei Fotos", "18 Foto mit Markierungen",
                    "19 Praxisfall", "20 Zitat", "21 Wissenscheck", "21 Wissenscheck", "22 Zusammenfassung",
                    "23 Abschluss"]


@pytest.fixture(scope="module")
def built(template_path, tmp_path_factory, m):
    out = tmp_path_factory.mktemp("deck") / "all.pptx"
    report = build_deck(copy.deepcopy(ALL_TYPES), template_path, out, SKILL / "examples", m)
    return out, Presentation(out), report


def test_layout_per_type(built):
    _, prs, report = built
    assert [s.slide_layout.name for s in prs.slides] == EXPECTED_LAYOUTS
    assert report.slides == len(EXPECTED_LAYOUTS)


def test_clean_build(built, m):
    out, _, _ = built
    assert check_deck(out, m) == []


def test_header_nav_marks_chapter(built):
    _, prs, _ = built
    names = [s.name for s in prs.slides[4].shapes]
    assert "nav-3-active" in names and "nav-underline" in names
    assert "nav-1" in names and "nav-4" in names


def test_chapter_ring(built):
    _, prs, _ = built
    names = [s.name for s in prs.slides[2].shapes if s.name.startswith("ring-seg")]
    assert names == ["ring-seg-1-done", "ring-seg-2-active", "ring-seg-3-todo", "ring-seg-4-todo"]


def test_footer_and_page_number(built):
    _, prs, _ = built
    shapes = {s.name: s for s in prs.slides[4].shapes}
    assert shapes["footer"].text_frame.text == "Zukunft im Kreislauf"
    assert shapes["page-number"].text_frame.text == "05"
    assert "footer" not in {s.name for s in prs.slides[0].shapes}


def test_legal_chips(built):
    _, prs, _ = built
    names = [s.name for s in prs.slides[4].shapes]
    assert "legal-label" in names and "legal-chip-1" in names


def test_notes(built):
    _, prs, _ = built
    assert prs.slides[4].notes_slide.notes_text_frame.text == "Sprechertext zu Folie fünf."


def test_briefs_reported(built):
    _, _, report = built
    assert (6, "Foto: Mitarbeiterin mit Tablet") in report.briefs
    assert len(report.briefs) == 5


def test_quote_without_photo_moves_left(built):
    _, prs, _ = built
    from campus.tokens import PH_IDX

    quote = next(ph for ph in prs.slides[19].placeholders if ph.placeholder_format.idx == PH_IDX["Quote"])
    assert quote.left == 120 * 6350


def test_optional_absent(template_path, tmp_path, m):
    plan = {"presentation": {"title": "Kurz"}, "slides": [{"type": "cover"}, {
        "type": "statement", "big": "3"}]}
    out = tmp_path / "kurz.pptx"
    build_deck(plan, template_path, out, tmp_path, m)
    prs = Presentation(out)
    assert check_deck(out, m) == []
    assert not any(s.name.startswith("nav-") for s in prs.slides[1].shapes)   # ohne Kapitel keine Navigation


SCRIPTS = SKILL / "scripts"


def test_cli_ok(tmp_path, template_path):
    plan = tmp_path / "plan.json"
    plan.write_text(json.dumps({"presentation": {"title": "T"}, "slides": [{"type": "cover"}]}), encoding="utf-8")
    out = tmp_path / "out.pptx"
    res = subprocess.run([sys.executable, SCRIPTS / "build_deck.py", plan, "--out", out, "--template", template_path],
                         capture_output=True, text=True)
    assert res.returncode == 0, res.stderr
    assert json.loads(res.stdout)["slides"] == 1 and out.exists()


def test_cli_plan_error(tmp_path):
    plan = tmp_path / "plan.json"
    plan.write_text(json.dumps({"presentation": {"title": "T"}, "slides": [{"type": "nope"}]}), encoding="utf-8")
    res = subprocess.run([sys.executable, SCRIPTS / "build_deck.py", plan, "--out", tmp_path / "x.pptx"],
                         capture_output=True, text=True)
    assert res.returncode == 2 and "unbekannter Typ" in res.stdout


def test_cli_missing_template(tmp_path):
    plan = tmp_path / "plan.json"
    plan.write_text(json.dumps({"presentation": {"title": "T"}, "slides": [{"type": "cover"}]}), encoding="utf-8")
    res = subprocess.run([sys.executable, SCRIPTS / "build_deck.py", plan, "--out", tmp_path / "x.pptx",
                          "--template", tmp_path / "fehlt.pptx"], capture_output=True, text=True)
    assert res.returncode == 2 and "Vorlage nicht gefunden" in res.stdout


def test_make_template_gallery(tmp_path):
    out = tmp_path / "vorlage.pptx"
    res = subprocess.run([sys.executable, SCRIPTS / "make_template.py", "--out", out,
                          "--gallery", SKILL / "examples" / "layout-gallery.json"], capture_output=True, text=True)
    assert res.returncode == 0, res.stderr
    prs = Presentation(out)
    assert len(prs.slide_layouts) == 23 and len(prs.slides) == 25
    assert (tmp_path / "layout-spec.json").exists()
