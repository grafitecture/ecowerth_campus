import pytest
from pptx import Presentation

from campus.layouts import LAYOUTS
from campus.template import layout_spec
from campus.tokens import THEME_MAJOR, THEME_MINOR

EXPECTED = ["01 Titel", "02 Agenda", "03 Kapitel", "04 Einstieg", "05 Punkte", "06 Punkte + Foto",
            "07 Punkte + Grafik", "08 Aussage", "09 Kennzahlen", "10 Marken", "11 Checkliste", "12 Schritte",
            "13 Vergleich", "14 Tabelle", "15 Skala", "16 Foto", "17 Zwei Fotos", "18 Foto mit Markierungen",
            "19 Praxisfall", "20 Zitat", "21 Wissenscheck", "22 Zusammenfassung", "23 Abschluss"]


@pytest.fixture(scope="module")
def prs(template_path):
    return Presentation(str(template_path))


def test_slide_size(prs):
    assert (prs.slide_width, prs.slide_height) == (12192000, 6858000)


def test_layout_names(prs):
    assert [lay.name for lay in prs.slide_layouts] == EXPECTED


def test_keys_unique():
    keys = [ldef.key for ldef in LAYOUTS]
    assert len(keys) == len(set(keys)) == 23


def test_theme_fonts(prs):
    theme = prs.slide_master.part.part_related_by(
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/theme").blob.decode("utf-8")
    assert f'typeface="{THEME_MAJOR}"' in theme and f'typeface="{THEME_MINOR}"' in theme
    assert "6EA52E" in theme and "043751" in theme


def _names(layout):
    return [s.name for s in layout.shapes]


def test_header_and_footer_lines(prs):
    for layout in prs.slide_layouts:
        frame = next(ldef.frame for ldef in LAYOUTS if ldef.name == layout.name)
        names = _names(layout)
        if frame in ("light", "dark"):
            assert {"head-logo", "head-line", "foot-line"} <= set(names), layout.name
        else:
            assert "head-line" not in names, layout.name


def test_cover_carries_the_loop(prs):
    names = _names(prs.slide_layouts[0])
    assert {"cover-logo", "cover-loop-hair", "cover-loop-core", "cover-loop-dot"} <= set(names)
    assert [n for n in names if n.startswith("cover-loop-arc")] == ["cover-loop-arc-1", "cover-loop-arc-2",
                                                                     "cover-loop-arc-3"]
    assert [n for n in names if n.startswith("cover-loop-name")] == [f"cover-loop-name-{i}" for i in (1, 2, 3)]


def test_facts_band_in_layout(prs):
    assert "facts-band" in _names(prs.slide_layouts[8])


def test_placeholders_match_spec(prs):
    spec = layout_spec()
    for layout in prs.slide_layouts:
        expected = spec[layout.name]["placeholders"]
        found = {ph.name for ph in layout.placeholders}
        assert found == set(expected), layout.name


def test_spec_json(tmp_path):
    from campus.template import write_layout_spec
    import json

    path = tmp_path / "layout-spec.json"
    write_layout_spec(path)
    data = json.loads(path.read_text(encoding="utf-8"))
    assert list(data) == EXPECTED
    assert data["05 Punkte"]["placeholders"]["Points"]["levels"][0]["font"] == "display"
