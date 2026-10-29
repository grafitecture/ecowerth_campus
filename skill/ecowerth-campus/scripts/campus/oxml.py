"""XML-Ergänzungen, die python-pptx nicht direkt anbietet.

Alle Änderungen laufen über lxml auf `shape._element` und halten die
Schema-Reihenfolge von DrawingML ein.
"""
from typing import Literal

from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from .tokens import FONTS, px

_GEOM = ("a:xfrm", "a:custGeom", "a:prstGeom")
_FILLS = ("a:noFill", "a:solidFill", "a:gradFill", "a:blipFill", "a:pattFill", "a:grpFill")


def _clr(color: str, alpha_pct: int | None = None):
    clr = OxmlElement("a:srgbClr")
    clr.set("val", color)
    if alpha_pct is not None:
        alpha = OxmlElement("a:alpha")
        alpha.set("val", str(int(alpha_pct * 1000)))
        clr.append(alpha)
    return clr


def _insert_after(parent, child, predecessors: tuple[str, ...]) -> None:
    """Fügt `child` hinter dem letzten vorhandenen Vorgänger ein (sonst an den Anfang)."""
    tags = {qn(t) for t in predecessors}
    last = None
    for el in parent:
        if el.tag in tags:
            last = el
    if last is None:
        parent.insert(0, child)
    else:
        last.addnext(child)


def _set_alpha(clr, alpha_pct: int) -> None:
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str(int(alpha_pct * 1000)))
    clr.append(alpha)


def set_fill_alpha(shape, alpha_pct: int) -> None:
    clr = shape._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
    _set_alpha(clr, alpha_pct)


def set_line_alpha(shape, alpha_pct: int) -> None:
    clr = shape._element.spPr.find(qn("a:ln")).find(qn("a:solidFill")).find(qn("a:srgbClr"))
    _set_alpha(clr, alpha_pct)


def set_line_cap(shape, cap: Literal["rnd", "sq", "flat"] = "rnd") -> None:
    shape._element.spPr.find(qn("a:ln")).set("cap", cap)


def set_linear_fill(shape, stops: list[tuple[int, str]], angle_deg: float) -> None:
    """Linearer Verlauf (Website: linear-gradient(140deg, …)); Positionen in Prozent."""
    spPr = shape._element.spPr
    for tag in _FILLS:
        for old in spPr.findall(qn(tag)):
            spPr.remove(old)
    grad = OxmlElement("a:gradFill")
    grad.set("rotWithShape", "1")
    gs_lst = OxmlElement("a:gsLst")
    for pos, color in stops:
        gs = OxmlElement("a:gs")
        gs.set("pos", str(int(pos * 1000)))
        gs.append(_clr(color))
        gs_lst.append(gs)
    grad.append(gs_lst)
    lin = OxmlElement("a:lin")
    lin.set("ang", str(int(angle_deg * 60000) % 21600000))
    lin.set("scaled", "0")
    grad.append(lin)
    _insert_after(spPr, grad, _GEOM)


def set_arc_angles(shape, start_deg: float, end_deg: float) -> None:
    """Winkel einer Bogenform (Preset „arc“): im Uhrzeigersinn ab 3 Uhr, in Grad."""
    prst = shape._element.spPr.find(qn("a:prstGeom"))
    av = prst.find(qn("a:avLst"))
    if av is None:
        av = OxmlElement("a:avLst")
        prst.append(av)
    for old in list(av):
        av.remove(old)
    for name, deg in (("adj1", start_deg), ("adj2", end_deg)):
        gd = OxmlElement("a:gd")
        gd.set("name", name)
        gd.set("fmla", f"val {int(round(deg * 60000)) % 21600000}")
        av.append(gd)


def _font(style: dict) -> tuple[str, bool]:
    return FONTS[style.get("font", "body")]


def _def_rpr(style: dict):
    typeface, bold = _font(style)
    d = OxmlElement("a:defRPr")
    d.set("sz", str(int(round(style["size_pt"] * 100))))
    d.set("b", "1" if bold else "0")
    d.set("i", "0")
    if style.get("caps"):
        d.set("cap", "all")
    if style.get("spc"):
        d.set("spc", str(int(style["spc"])))
    fill = OxmlElement("a:solidFill")
    fill.append(_clr(style["color"]))
    d.append(fill)
    for tag in ("a:latin", "a:ea", "a:cs"):
        font = OxmlElement(tag)
        font.set("typeface", typeface)
        d.append(font)
    return d


def _spacing(tag: str, pts: float):
    el = OxmlElement(tag)
    spc = OxmlElement("a:spcPts")
    spc.set("val", str(int(round(pts * 100))))
    el.append(spc)
    return el


def level_ppr(level: int, style: dict):
    """`a:lvlNpPr` (N = level + 1) aus einem Ebenen-Stil."""
    p = OxmlElement(f"a:lvl{level + 1}pPr")
    p.set("marL", str(px(style.get("mar_l_px", 0))))
    p.set("indent", str(px(style.get("indent_px", 0))))
    p.set("algn", {"left": "l", "center": "ctr", "right": "r"}[style.get("align", "left")])
    if style.get("ln_spc_pct"):
        ln = OxmlElement("a:lnSpc")
        pct = OxmlElement("a:spcPct")
        pct.set("val", str(int(style["ln_spc_pct"] * 1000)))
        ln.append(pct)
        p.append(ln)
    p.append(_spacing("a:spcBef", style.get("spc_before_pt", 0)))
    p.append(_spacing("a:spcAft", style.get("spc_after_pt", 0)))
    p.append(OxmlElement("a:buNone"))
    if style.get("tab_px"):
        tabs = OxmlElement("a:tabLst")
        tab = OxmlElement("a:tab")
        tab.set("pos", str(px(style["tab_px"])))
        tab.set("algn", "l")
        tabs.append(tab)
        p.append(tabs)
    p.append(_def_rpr(style))
    return p


def to_placeholder(
    shape,
    ph_type: Literal["title", "body", "pic"],
    idx: int | None,
    name: str,
    levels: list[dict],
    prompt: list[str],
    anchor: str = "t",
    insets_px=(0, 0, 0, 0),
    geom: str = "rect",
) -> None:
    sp = shape._element
    sp.nvSpPr.cNvPr.set("name", name)

    cNvSpPr = sp.nvSpPr.cNvSpPr
    cNvSpPr.attrib.pop("txBox", None)
    for old in cNvSpPr.findall(qn("a:spLocks")):
        cNvSpPr.remove(old)
    locks = OxmlElement("a:spLocks")
    locks.set("noGrp", "1")
    cNvSpPr.insert(0, locks)

    nvPr = sp.nvSpPr.nvPr
    for old in nvPr.findall(qn("p:ph")):
        nvPr.remove(old)
    ph = OxmlElement("p:ph")
    ph.set("type", ph_type)
    if idx is not None and ph_type != "title":
        ph.set("idx", str(idx))
    ph.set("hasCustomPrompt", "1")
    nvPr.insert(0, ph)

    style_el = sp.find(qn("p:style"))
    if style_el is not None:
        sp.remove(style_el)

    spPr = sp.spPr
    prst = spPr.find(qn("a:prstGeom"))
    if prst is None:
        prst = OxmlElement("a:prstGeom")
        _insert_after(spPr, prst, ("a:xfrm",))
    prst.set("prst", geom)
    if prst.find(qn("a:avLst")) is None:
        prst.append(OxmlElement("a:avLst"))

    txBody = sp.find(qn("p:txBody"))
    if txBody is None:
        txBody = OxmlElement("p:txBody")
        sp.append(txBody)
    for child in list(txBody):
        txBody.remove(child)

    body_pr = OxmlElement("a:bodyPr")
    body_pr.set("wrap", "square")
    for attr, value in zip(("lIns", "tIns", "rIns", "bIns"), insets_px):
        body_pr.set(attr, str(px(value)))
    body_pr.set("anchor", {"t": "t", "ctr": "ctr", "b": "b"}[anchor])
    body_pr.append(OxmlElement("a:noAutofit"))
    txBody.append(body_pr)

    lst = OxmlElement("a:lstStyle")
    for level, style in enumerate(levels):
        lst.append(level_ppr(level, style))
    txBody.append(lst)

    for level, text in enumerate(prompt):
        p = OxmlElement("a:p")
        if level:
            ppr = OxmlElement("a:pPr")
            ppr.set("lvl", str(level))
            p.append(ppr)
        r = OxmlElement("a:r")
        rpr = OxmlElement("a:rPr")
        rpr.set("lang", "de-DE")
        rpr.set("dirty", "0")
        r.append(rpr)
        t = OxmlElement("a:t")
        t.text = text
        r.append(t)
        p.append(r)
        txBody.append(p)


_ALIGN = {"left": "LEFT", "center": "CENTER", "right": "RIGHT"}
_ANCHOR = {"t": "TOP", "ctr": "MIDDLE", "b": "BOTTOM"}


def style_run(run, style: dict) -> None:
    """Setzt Schrift, Größe, Farbe, Versalien und Laufweite eines Laufs explizit."""
    from pptx.dml.color import RGBColor
    from pptx.util import Pt

    typeface, bold = _font(style)
    run.font.name = typeface
    run.font.size = Pt(style["size_pt"])
    run.font.bold = bold
    run.font.italic = False
    run.font.color.rgb = RGBColor.from_string(style["color"])
    rpr = run._r.get_or_add_rPr()
    rpr.set("lang", "de-DE")
    if style.get("caps"):
        rpr.set("cap", "all")
    if style.get("spc"):
        rpr.set("spc", str(int(style["spc"])))


def fill_text_frame(tf, paragraphs: list[tuple[str, dict]], anchor: str = "t") -> None:
    """Füllt einen Textrahmen: ein Lauf pro Absatz, ohne Innenabstand, ohne Autofit."""
    from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
    from pptx.util import Pt

    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = getattr(MSO_ANCHOR, _ANCHOR[anchor])
    for i, (text, style) in enumerate(paragraphs):
        para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        para.alignment = getattr(PP_ALIGN, _ALIGN[style.get("align", "left")])
        if style.get("ln_spc_pct"):
            para.line_spacing = style["ln_spc_pct"] / 100
        if style.get("spc_before_pt"):
            para.space_before = Pt(style["spc_before_pt"])
        if style.get("spc_after_pt"):
            para.space_after = Pt(style["spc_after_pt"])
        run = para.add_run()
        run.text = text
        style_run(run, style)
