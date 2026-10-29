"""Rahmen-Widgets: Kapitel-Navigation in der Kopfzeile, Fußzeile, Quelle und Paragrafen-Chips,
Fokus-Box, Fotos und Bildbriefings."""
import tempfile
from pathlib import Path

from PIL import Image, ImageOps
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement

from . import oxml
from .draw import hline, rect, shape, style, textbox
from .measure import Measurer
from .tokens import (CONTENT_W, EMU_PER_PX, FONTS, FOOT_TEXT_Y, HEAD_H, KICKER_SPC, LABEL_SPC, LOGO_BOX, MARGIN_X,
                     NAV_SPC, SIZE_PT, SOURCE_Y, px)

ASSETS = Path(__file__).resolve().parents[2] / "assets"

NAV_GAP = 40           # Abstand zwischen den Kapiteln in der Kopfzeile
NAV_NUM_GAP = 10       # Abstand Nummer → Name
NAV_LEFT = LOGO_BOX[0] + LOGO_BOX[2] + 120   # die Navigation bleibt rechts vom Logo
UNDERLINE_H = 4


def _runs_box(slide, name: str, box, runs: list[tuple[str, dict]], align: str = "left", anchor: str = "ctr"):
    """Ein Absatz mit mehreren Läufen (z. B. Nummer in Mono, Name in Plex Sans)."""
    x, y, w, h = box
    tb = slide.shapes.add_textbox(px(x), px(y), px(w), px(h))
    tb.name = name
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "ctr": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
    para = tf.paragraphs[0]
    para.alignment = {"left": PP_ALIGN.LEFT, "right": PP_ALIGN.RIGHT, "center": PP_ALIGN.CENTER}[align]
    for text, st in runs:
        run = para.add_run()
        run.text = text
        oxml.style_run(run, st)
    return tb


# ---------------------------------------------------------------- Kopfzeile

def nav_items(chapters: list[str], m: Measurer) -> list[dict]:
    """Breite je Kapitel-Eintrag; zu breite Navigation zeigt bei inaktiven Kapiteln nur die Nummer."""
    num_st = style("nav", "muted", "mono")
    name_st = style("nav", "muted", "body_medium", caps=True, spc=NAV_SPC)
    items = []
    space = m.width_px(" ", num_st["size_pt"], font=FONTS["mono"])
    for i, name in enumerate(chapters, 1):
        num = f"{i:02d}"
        wn = m.width_px(num, num_st["size_pt"], font=FONTS["mono"])
        wl = m.width_px(name.upper(), name_st["size_pt"], NAV_SPC / 100, FONTS["body_medium"])
        items.append({"n": i, "num": num, "name": name, "w_num": wn, "w_full": wn + space + wl})
    return items


def header_nav(slide, chapters: list[str], current: int | None, dark: bool, m: Measurer) -> None:
    """Kapitel rechts in der Kopfzeile, wie die Navigation der Website; das aktive ist unterstrichen."""
    if not chapters:
        return
    items = nav_items(chapters, m)
    room = MARGIN_X + CONTENT_W - NAV_LEFT
    compact = sum(it["w_full"] for it in items) + NAV_GAP * (len(items) - 1) > room
    widths = [it["w_full"] if not compact or it["n"] == current else it["w_num"] for it in items]
    x = MARGIN_X + CONTENT_W - sum(widths) - NAV_GAP * (len(items) - 1)
    base, active_color = ("muted_dark", "white") if dark else ("muted", "navy")
    _, y, _, h = LOGO_BOX
    for it, w in zip(items, widths):
        on = it["n"] == current
        color = active_color if on else base
        runs = [(it["num"], style("nav", "green_light" if (on and dark) else ("green_ink" if on else color), "mono"))]
        if w > it["w_num"] + 1:
            runs.append((" ", style("nav", color, "mono")))
            runs.append((it["name"], style("nav", color, "body_medium", caps=True, spc=NAV_SPC)))
        name = f"nav-{it['n']}" + ("-active" if on else "")
        _runs_box(slide, name, (x - 2, y, w + 24, h), runs)   # etwas Luft rechts: kein Umbruch durch Rundung
        if on:
            rect(slide.shapes, "nav-underline", (x, HEAD_H - UNDERLINE_H, w, UNDERLINE_H),
                 "green_light" if dark else "green")
        x += w + NAV_GAP


# ---------------------------------------------------------------- Fußzeile, Quelle, Paragrafen-Chips

def footer(slide, text: str, number: int, dark: bool) -> None:
    color = "muted_dark" if dark else "muted"
    textbox(slide.shapes, "footer", (MARGIN_X, FOOT_TEXT_Y, 1300, 30),
            [(text, style("footer", color, "mono_medium", caps=True, spc=LABEL_SPC))], anchor="ctr")
    textbox(slide.shapes, "page-number", (MARGIN_X + CONTENT_W - 200, FOOT_TEXT_Y, 200, 30),
            [(f"{number:02d}", style("footer", color, "mono_medium", align="right"))], anchor="ctr")


CHIP_H = 34
CHIP_PAD = 12
CHIP_GAP = 10


def chips(slide, items: list[str], x: float, y: float, m: Measurer, dark: bool = False,
          prefix: str = "chip") -> float:
    """Paragrafen-Chips wie auf der Website (IBM Plex Mono auf hellblauem Grund). Gibt das rechte Ende zurück."""
    st = style("chip", "blue_ink", "mono", align="center")
    for i, item in enumerate(items, 1):
        w = m.width_px(item, st["size_pt"], font=FONTS["mono"]) + 2 * CHIP_PAD
        chip = shape(slide.shapes, MSO_SHAPE.ROUNDED_RECTANGLE, f"{prefix}-{i}", (x, y, w, CHIP_H),
                     fill="blue_tint" if not dark else "line_dark", radius_px=4)
        oxml.fill_text_frame(chip.text_frame, [(item, {**st, "color": "0F6AA6" if not dark else "E8EEF4"})],
                             anchor="ctr")
        x += w + CHIP_GAP
    return x - CHIP_GAP


def source_row(slide, source: str | None, legal: list[str] | None, dark: bool, m: Measurer) -> None:
    """Rechtsgrundlagen als Chips und/oder Quelle in einer Zeile über der Fußzeile."""
    x = MARGIN_X
    if legal:
        label = "Rechtsgrundlage"
        st = style("label", "green_light" if dark else "green_ink", "mono_medium", caps=True, spc=LABEL_SPC)
        w = m.width_px(label.upper(), st["size_pt"], LABEL_SPC / 100, FONTS["mono_medium"]) + 4
        textbox(slide.shapes, "legal-label", (x, SOURCE_Y, w + 20, CHIP_H), [(label, st)], anchor="ctr")
        x = chips(slide, legal, x + w + 16, SOURCE_Y, m, dark, prefix="legal-chip") + 32
    if source:
        label_st = style("label", "green_light" if dark else "green_ink", "mono_medium", caps=True, spc=LABEL_SPC)
        text_st = style("small", "muted_dark" if dark else "muted")
        _runs_box(slide, "source", (x, SOURCE_Y, MARGIN_X + CONTENT_W - x, CHIP_H),
                  [("Quelle", label_st), ("   ", text_st), (source, text_st)])


# ---------------------------------------------------------------- Fokus-Box

def focus_box(slide, box: tuple[float, float, float], label: str, title: str, text: str, m: Measurer) -> float:
    """Dunkelblaue Box mit grüner Kante: der eine Satz, der hängen bleiben soll."""
    x, y, w = box
    pad_x, pad_y, inner = 48, 44, w - 96
    label_st = style("label", "green_light", "mono_medium", caps=True, spc=LABEL_SPC)
    title_st = style("focus_title", "white", "display")
    text_st = style("body", "on_dark")
    label_h = m.line_height_px(label_st["size_pt"], font=FONTS["mono_medium"])
    title_h = m.text_height_px(title, title_st, inner)
    text_h = m.text_height_px(text, text_st, inner) if text else 0
    height = round(pad_y + label_h + 12 + title_h + (18 + text_h if text else 0) + pad_y)
    rect(slide.shapes, "focus-card", (x, y, w, height), "navy")
    rect(slide.shapes, "focus-accent", (x, y, 6, height), "green")
    top = y + pad_y
    textbox(slide.shapes, "focus-label", (x + pad_x, top, inner, label_h + 2), [(label, label_st)])
    top += label_h + 12
    textbox(slide.shapes, "focus-title", (x + pad_x, top, inner, title_h + 2), [(title, title_st)])
    if text:
        textbox(slide.shapes, "focus-text", (x + pad_x, top + title_h + 18, inner, text_h + 2), [(text, text_st)])
    return height


# ---------------------------------------------------------------- Fotos

def _as_px(ph) -> tuple[float, float, float, float]:
    return tuple(v / EMU_PER_PX for v in (ph.left, ph.top, ph.width, ph.height))


MIN_PHOTO_RATIO = 0.5  # unter halber Sollgröße gilt ein Foto als zu klein (Platzhalter + Fotoliste)
MAX_PHOTO_RATIO = 2.0  # mehr als doppelte Auflösung bringt nichts und bläht das Deck auf


def _ratio(size: tuple[int, int], box: tuple[float, float, float, float], fit: bool) -> float:
    """Bildpixel je Folienpixel: beim Füllen (beschnitten) zählt die knappere Seite, beim Einpassen die größere."""
    rx, ry = size[0] / box[2], size[1] / box[3]
    return max(rx, ry) if fit else min(rx, ry)


def _usable_image(path: Path, box, fit: bool = False) -> tuple[Path, tuple[int, int]]:
    """Bild am Inhalt erkennen, EXIF-Drehung anwenden, auf höchstens doppelte Rahmenauflösung verkleinern;
    alles außer unverändertem JPEG/PNG wird neu gespeichert. Gibt Pfad und Pixelmaß zurück."""
    with Image.open(path) as im:
        im.load()
        rotated = im.getexif().get(0x0112, 1) != 1
        img = ImageOps.exif_transpose(im) if rotated else im
        scale = MAX_PHOTO_RATIO / _ratio(img.size, box, fit)
        if im.format in ("JPEG", "PNG") and not rotated and scale >= 1:
            return path, im.size
        if scale < 1:
            img = img.resize((max(1, round(img.width * scale)), max(1, round(img.height * scale))), Image.LANCZOS)
        alpha = img.mode in ("RGBA", "LA") or (img.mode == "P" and "transparency" in img.info)
        out = Path(tempfile.mkdtemp(prefix="campus-img-")) / (path.stem + (".png" if alpha else ".jpg"))
        if alpha:
            img.convert("RGBA").save(out, "PNG", optimize=True)
        else:
            img.convert("RGB").save(out, "JPEG", quality=88, optimize=True)
        return out, img.size


def _prepare(photo: dict | None, base_dir: Path, box, warnings: list[str]):
    """Foto prüfen: gibt (Pfad, Pixelmaß) zurück oder den Briefing-Text für eine Platzhalterkarte."""
    photo = photo or {}
    file, brief = photo.get("file"), photo.get("brief")
    path = Path(file) if file else None
    if path is not None and not path.is_absolute():
        path = Path(base_dir) / path
    if path is None or not path.exists():
        if file:
            warnings.append(f"Foto fehlt: {file}")
        return brief or f"Foto fehlt: {file}"
    fit = bool(photo.get("fit"))
    try:
        usable, size = _usable_image(path, box, fit)
    except Exception as err:  # beschädigt, abgeschnitten, unbekanntes Format: Briefing statt Absturz
        warnings.append(f"Foto unlesbar: {file} ({type(err).__name__})")
        return brief or f"Foto unlesbar: {file}"
    if _ratio(size, box, fit) < MIN_PHOTO_RATIO:
        warnings.append(f"Foto zu klein: {file} ({size[0]}×{size[1]} px für {box[2]:.0f}×{box[3]:.0f} px)")
        return brief or f"Foto zu klein: {file}"
    return usable, size


def _fitted(size: tuple[int, int], box) -> tuple[float, float, float, float]:
    x, y, w, h = box
    scale = min(w / size[0], h / size[1])
    fw, fh = size[0] * scale, size[1] * scale
    return x + (w - fw) / 2, y + (h - fh) / 2, fw, fh


def place_photo(slide, idx: int, photo: dict | None, base_dir: Path, warnings: list[str]) -> str | None:
    ph = next(p for p in slide.placeholders if p.placeholder_format.idx == idx)
    box = _as_px(ph)
    prepared = _prepare(photo, base_dir, box, warnings)
    if isinstance(prepared, str):
        return _brief_in_place(slide, ph, prepared)
    usable, size = prepared
    pic = ph.insert_picture(str(usable))
    # Die graue Leerfläche des Platzhalters nicht erben: transparente Bilder zeigen den Folienhintergrund.
    sp_pr = pic._element.spPr
    sp_pr.insert(1 + sum(1 for c in sp_pr if c.tag in (qn("a:xfrm"), qn("a:prstGeom"), qn("a:custGeom"))),
                 OxmlElement("a:noFill"))
    if (photo or {}).get("fit"):  # ganz zeigen: ohne Beschnitt, mittig im Rahmen
        pic.crop_left = pic.crop_right = pic.crop_top = pic.crop_bottom = 0
        x, y, w, h = _fitted(size, box)
        pic.left, pic.top, pic.width, pic.height = px(x), px(y), px(w), px(h)
    return None


def add_photo(slide, name: str, box, photo: dict | None, base_dir: Path, warnings: list[str]) -> str | None:
    """Foto als freies Bild in einen Rahmen setzen (gefüllt und mittig beschnitten oder eingepasst)."""
    prepared = _prepare(photo, base_dir, box, warnings)
    if isinstance(prepared, str):
        brief_card(slide, box, prepared)
        return prepared
    usable, size = prepared
    fit = bool((photo or {}).get("fit"))
    x, y, w, h = _fitted(size, box) if fit else box
    pic = slide.shapes.add_picture(str(usable), px(x), px(y), px(w), px(h))
    pic.name = name
    if not fit:
        image_aspect, box_aspect = size[0] / size[1], w / h
        if image_aspect > box_aspect:
            pic.crop_left = pic.crop_right = (1 - box_aspect / image_aspect) / 2
        else:
            pic.crop_top = pic.crop_bottom = (1 - image_aspect / box_aspect) / 2
    return None


def _brief_in_place(slide, ph, text: str) -> str:
    box = _as_px(ph)
    tree = ph._element.getparent()
    anchor = list(tree).index(ph._element)
    tree.remove(ph._element)
    count = len(tree)
    brief_card(slide, box, text)
    # Die Karte nimmt die Ebene des Platzhalters ein, damit Texte darüber sichtbar bleiben.
    for offset, el in enumerate(list(tree)[count:]):
        tree.remove(el)
        tree.insert(anchor + offset, el)
    return text


def brief_card(slide, box, brief: str) -> None:
    """Bildplatzhalter wie auf der Website: Verlauf Dunkelblau → Blau, Briefing in Mono unten links."""
    x, y, w, h = box
    card = rect(slide.shapes, "brief-card", box, "navy")
    oxml.set_linear_fill(card, [(0, "043751"), (90, "117EC2")], 40)
    pad = 28 if w > 400 else 18
    label_st = style("label", "green_light", "mono_medium", caps=True, spc=LABEL_SPC)
    text_st = style("small", "white", "mono", spc_before_pt=4)
    box_h = min(h - 2 * pad, 150)
    textbox(slide.shapes, "brief-text", (x + pad, y + h - pad - box_h, w - 2 * pad, box_h),
            [("Platzhalter", label_st), (brief, text_st)], anchor="b")
