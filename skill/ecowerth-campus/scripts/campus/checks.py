"""Prüfungen für gebaute Decks: Überlauf, Ränder, leere Platzhalter, fremde Schriften, fehlende Zeichen."""
from dataclasses import dataclass
from pathlib import Path
from typing import Literal

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn
from pptx.shapes.placeholder import PlaceholderPicture

from .measure import DEFAULT_FONT, Measurer, Para
from .tokens import EMU_PER_PX, FONTS, PH_IDX, THEME_MAJOR, THEME_MINOR

IDX_NAME = {idx: name for name, idx in PH_IDX.items()}
ALLOWED_FONTS = {None, "+mj-lt", "+mn-lt", THEME_MAJOR, THEME_MINOR} | {name for name, _ in FONTS.values()}
THEME_FONT = {"+mj-lt": THEME_MAJOR, "+mn-lt": THEME_MINOR}


@dataclass
class Issue:
    slide: int
    kind: Literal["overflow", "empty", "font", "glyph", "word", "bounds"]
    shape: str
    detail: str


def _px(emu) -> float:
    return (emu or 0) / EMU_PER_PX


def _logical_name(shape) -> str:
    if shape.is_placeholder:
        return IDX_NAME.get(shape.placeholder_format.idx, shape.name)
    return shape.name


def _layout_ph(slide, idx: int):
    return next((p for p in slide.slide_layout.placeholders if p.placeholder_format.idx == idx), None)


def _level_props(layout_ph, level: int) -> dict:
    """Effektive Absatz- und Zeichenwerte einer Ebene aus dem lstStyle des Layout-Platzhalters."""
    props = {"size_pt": None, "ln_pct": 100, "before": 0.0, "after": 0.0, "mar_l": 0.0, "spc": 0.0,
             "typeface": None, "bold": False}
    if layout_ph is None or layout_ph._element.txBody is None:
        return props
    lst = layout_ph._element.txBody.find(qn("a:lstStyle"))
    if lst is None:
        return props
    levels = [lst.find(qn(f"a:lvl{i}pPr")) for i in range(1, 10)]
    levels = [lv for lv in levels if lv is not None]
    if not levels:
        return props
    ppr = levels[min(level, len(levels) - 1)]
    props["mar_l"] = _px(int(ppr.get("marL", 0)))
    ln = ppr.find(qn("a:lnSpc"))
    if ln is not None and ln.find(qn("a:spcPct")) is not None:
        props["ln_pct"] = int(ln.find(qn("a:spcPct")).get("val")) / 1000
    for key, tag in (("before", "a:spcBef"), ("after", "a:spcAft")):
        el = ppr.find(qn(tag))
        if el is not None and el.find(qn("a:spcPts")) is not None:
            props[key] = int(el.find(qn("a:spcPts")).get("val")) / 100
    d = ppr.find(qn("a:defRPr"))
    if d is not None:
        if d.get("sz"):
            props["size_pt"] = int(d.get("sz")) / 100
        if d.get("spc"):
            props["spc"] = int(d.get("spc")) / 100
        props["bold"] = d.get("b") == "1"
        latin = d.find(qn("a:latin"))
        if latin is not None:
            props["typeface"] = latin.get("typeface")
    return props


def _insets(shape, layout_ph) -> tuple[float, float, float, float]:
    body = shape.text_frame._txBody.find(qn("a:bodyPr"))
    source = body
    if layout_ph is not None and not any(body.get(k) for k in ("lIns", "tIns", "rIns", "bIns")):
        source = layout_ph._element.txBody.find(qn("a:bodyPr"))
    defaults = {"lIns": 91440, "tIns": 45720, "rIns": 91440, "bIns": 45720}
    return tuple(_px(int(source.get(k, defaults[k]))) for k in ("lIns", "tIns", "rIns", "bIns"))


def _run_font(run, props: dict) -> tuple[str, bool]:
    typeface = run.font.name or props["typeface"]
    typeface = THEME_FONT.get(typeface, typeface)
    bold = run.font.bold if run.font.bold is not None else props["bold"]
    return (typeface or DEFAULT_FONT[0], bool(bold))


def _paras(shape, layout_ph) -> list[Para]:
    paras = []
    for p in shape.text_frame.paragraphs:
        props = _level_props(layout_ph, p.level)
        runs = [r for r in p.runs if r.text.strip()] or list(p.runs)
        # Der längste Lauf bestimmt Schrift und Größe des Absatzes (Nummern-Läufe sind kurz)
        main = max(runs, key=lambda r: len(r.text), default=None)
        size = (main.font.size.pt if main is not None and main.font.size is not None else None) \
            or props["size_pt"] or 18
        spc = props["spc"]
        if main is not None and main._r.rPr is not None and main._r.rPr.get("spc"):
            spc = int(main._r.rPr.get("spc")) / 100
        ln_pct = props["ln_pct"]
        if isinstance(p.line_spacing, float):
            ln_pct = p.line_spacing * 100
        before = p.space_before.pt if p.space_before is not None else props["before"]
        after = p.space_after.pt if p.space_after is not None else props["after"]
        font = _run_font(main, props) if main is not None else (props["typeface"] or DEFAULT_FONT[0], props["bold"])
        text = p.text.upper() if _caps(main, layout_ph, p.level) else p.text
        paras.append(Para(text, size, ln_pct, before, after, spc, props["mar_l"], font))
    return paras


def _caps(run, layout_ph, level: int) -> bool:
    if run is not None and run._r.rPr is not None and run._r.rPr.get("cap") == "all":
        return True
    if layout_ph is None or layout_ph._element.txBody is None:
        return False
    lst = layout_ph._element.txBody.find(qn("a:lstStyle"))
    ppr = lst.find(qn(f"a:lvl{level + 1}pPr")) if lst is not None else None
    d = ppr.find(qn("a:defRPr")) if ppr is not None else None
    return d is not None and d.get("cap") == "all"


def check_deck(pptx: Path, m: Measurer) -> list[Issue]:
    prs = Presentation(str(pptx))
    issues: list[Issue] = []
    slide_w, slide_h = _px(prs.slide_width), _px(prs.slide_height)
    for n, slide in enumerate(prs.slides, 1):
        for shape in slide.shapes:
            name = _logical_name(shape)
            layout_ph = _layout_ph(slide, shape.placeholder_format.idx) if shape.is_placeholder else None
            x, y = _px(shape.left), _px(shape.top)
            if x < -1 or y < -1 or x + _px(shape.width) > slide_w + 1 or y + _px(shape.height) > slide_h + 1:
                issues.append(Issue(n, "bounds", name, "ragt über den Folienrand"))
            if isinstance(shape, PlaceholderPicture) or shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                continue
            if not shape.has_text_frame:
                continue
            text = shape.text_frame.text
            if not text.strip():
                if shape.is_placeholder:
                    issues.append(Issue(n, "empty", name, "Platzhalter ohne Inhalt"))
                continue
            for p in shape.text_frame.paragraphs:
                props = _level_props(layout_ph, p.level)
                for r in p.runs:
                    if r.font.name not in ALLOWED_FONTS:
                        issues.append(Issue(n, "font", name, f"Schrift '{r.font.name}'"))
                    cmap = m.cmap(_run_font(r, props))
                    if cmap is not None:
                        missing = sorted({ch for ch in r.text if not ch.isspace() and ord(ch) not in cmap})
                        if missing:
                            issues.append(Issue(n, "glyph", name, f"fehlt in {_run_font(r, props)[0]}: "
                                                + " ".join(missing)))
            left, top, right, bottom = _insets(shape, layout_ph)
            width = _px(shape.width) - left - right
            height = _px(shape.height) - top - bottom
            paras = _paras(shape, layout_ph)
            for p in paras:
                for word in m.long_words(p.text.rsplit("\t", 1)[-1], p.size_pt, width - p.indent_px, p.spc_pt, p.font):
                    issues.append(Issue(n, "word", name, f"'{word}' ist breiter als der Rahmen – bricht mitten im Wort"))
            measured = m.block_height_px(paras, width)
            if measured > height + 2:
                issues.append(Issue(n, "overflow", name, f"{measured:.0f}px > {height:.0f}px"))
    return issues
