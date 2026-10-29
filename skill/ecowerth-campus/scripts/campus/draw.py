"""Zeichen-Grundlagen für Layouts und Folien: Textrahmen, Formen, Haarlinien, Bögen, Bilder.

Alle Funktionen nehmen eine Formen-Sammlung (`slide.shapes` oder die eines Layouts),
damit Rahmen-Elemente in der Vorlage und auf den Folien gleich entstehen.
Maße in Rasterpixeln (1920 × 1080), Farben als COLOR-Schlüssel oder Hex-Wert.
"""
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn

from . import oxml
from .tokens import COLOR, LINE_PX, SIZE_PT, px


def col(color: str) -> str:
    return COLOR.get(color, color)


def style(size: str | float, color: str, font: str = "body", **extra) -> dict:
    """Textstil: Größe als SIZE_PT-Schlüssel oder pt, Farbe als COLOR-Schlüssel oder Hex, Schriftrolle."""
    size_pt = SIZE_PT[size] if isinstance(size, str) else size
    return {"size_pt": size_pt, "color": col(color), "font": font, **extra}


def textbox(shapes, name: str, box, paragraphs: list[tuple[str, dict]], anchor: str = "t"):
    x, y, w, h = box
    tb = shapes.add_textbox(px(x), px(y), px(w), px(h))
    tb.name = name
    oxml.fill_text_frame(tb.text_frame, paragraphs, anchor=anchor)
    return tb


def _strip_style(sp) -> None:
    style_el = sp._element.find(qn("p:style"))
    if style_el is not None:
        sp._element.remove(style_el)


def apply_fill_line(sp, fill: str | None = None, line=None, fill_alpha: int | None = None) -> None:
    """Füllung und Kontur; line = (Farbe, Breite px[, gestrichelt][, Deckkraft %])."""
    if fill:
        sp.fill.solid()
        sp.fill.fore_color.rgb = RGBColor.from_string(col(fill))
        if fill_alpha is not None:
            oxml.set_fill_alpha(sp, fill_alpha)
    else:
        sp.fill.background()
    if line:
        color, width, *rest = line
        sp.line.color.rgb = RGBColor.from_string(col(color))
        sp.line.width = px(width)
        if rest and rest[0]:
            sp.line.dash_style = MSO_LINE.DASH
        if len(rest) > 1 and rest[1] is not None:
            oxml.set_line_alpha(sp, rest[1])
    else:
        sp.line.fill.background()


def shape(shapes, kind, name: str, box, fill: str | None = None, line=None, radius_px: float = 0,
          fill_alpha: int | None = None):
    x, y, w, h = box
    sp = shapes.add_shape(kind, px(x), px(y), px(w), px(h))
    sp.name = name
    _strip_style(sp)
    apply_fill_line(sp, fill, line, fill_alpha)
    if kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        sp.adjustments[0] = (radius_px or 4) / min(w, h)
    return sp


def rect(shapes, name: str, box, fill: str, **kw):
    return shape(shapes, MSO_SHAPE.RECTANGLE, name, box, fill=fill, **kw)


def hline(shapes, name: str, x: float, y: float, w: float, color: str = "grey", width: float = LINE_PX):
    """Waagerechte Haarlinie als schmale Fläche: steht in PowerPoint und LibreOffice gleich scharf."""
    return rect(shapes, name, (x, y - width / 2, w, width), color)


def vline(shapes, name: str, x: float, y: float, h: float, color: str = "grey", width: float = LINE_PX):
    return rect(shapes, name, (x - width / 2, y, width, h), color)


def circle(shapes, name: str, cx: float, cy: float, r: float, fill: str | None = None, line=None):
    return shape(shapes, MSO_SHAPE.OVAL, name, (cx - r, cy - r, 2 * r, 2 * r), fill=fill, line=line)


def arc(shapes, name: str, cx: float, cy: float, r: float, start_deg: float, end_deg: float, color: str,
        width_px: float):
    """Kreisbogen mit runden Enden (wie stroke-linecap: round), Winkel im Uhrzeigersinn ab 3 Uhr."""
    sp = shape(shapes, MSO_SHAPE.ARC, name, (cx - r, cy - r, 2 * r, 2 * r), line=(color, width_px))
    oxml.set_arc_angles(sp, start_deg, end_deg)
    oxml.set_line_cap(sp, "rnd")
    return sp


def picture(shapes, name: str, path, box):
    x, y, w, h = box
    pic = shapes.add_picture(str(path), px(x), px(y), px(w), px(h))
    pic.name = name
    return pic
