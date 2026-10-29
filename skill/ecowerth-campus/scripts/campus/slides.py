"""Ein Renderer pro Folientyp: Layout wählen, Platzhalter füllen, dynamische Elemente zeichnen."""
from dataclasses import dataclass, field
from pathlib import Path
from typing import Callable

from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.util import Pt

from . import motifs
from .layouts import CHAPTER_RING, LAYOUTS
from .measure import Measurer, Para
from .plan import chapters_of, default_kicker
from .tokens import EMU_PER_PX, FONTS, PH_IDX, px
from .widgets_content import (agenda_rows, bars_card, brand_columns, case_findings, checklist, closing_steps,
                              compare_columns, facts_band, markers, matrix_table, photo_steps, points_text,
                              quiz_options, scale_caption, scale_chips, steps_grid, summary_items, tag_cloud,
                              timeline)
from .widgets_frame import focus_box, footer, header_nav, place_photo, source_row

LAYOUT_BY_KEY = {ldef.key: ldef.name for ldef in LAYOUTS}
LAYOUT_BY_NAME = {ldef.name: ldef for ldef in LAYOUTS}
DARK_KEYS = {ldef.key for ldef in LAYOUTS if ldef.frame in ("dark", "cover", "full")}
NO_HEADER = {"cover", "photo"}
NO_FOOTER = {"cover", "chapter", "photo"}


@dataclass
class BuildReport:
    slides: int = 0
    briefs: list[tuple[int, str]] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)


@dataclass
class Ctx:
    prs: object
    plan: dict
    m: Measurer
    base_dir: Path
    report: BuildReport

    @property
    def pres(self) -> dict:
        return self.plan["presentation"]

    @property
    def chapters(self) -> list[str]:
        return chapters_of(self.plan)


# ---------------------------------------------------------------- Grundfunktionen

def new_slide(ctx: Ctx, layout_name: str):
    layout = next(lay for lay in ctx.prs.slide_layouts if lay.name == layout_name)
    slide = ctx.prs.slides.add_slide(layout)
    ctx.report.slides += 1
    return slide


def _placeholder(slide, name: str):
    for ph in slide.placeholders:
        if name == "Title":
            if ph.placeholder_format.type in (PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE):
                return ph
        elif ph.placeholder_format.idx == PH_IDX[name]:
            return ph
    return None


def fill(slide, name: str, text: str | None) -> None:
    ph = _placeholder(slide, name)
    if ph is None:
        return
    if not text or not str(text).strip():
        ph._element.getparent().remove(ph._element)
        return
    tf = ph.text_frame
    for extra in list(tf.paragraphs)[1:]:
        extra._p.getparent().remove(extra._p)
    lines = str(text).split("\n")
    tf.paragraphs[0].add_run().text = lines[0]
    for line in lines[1:]:
        tf.add_paragraph().add_run().text = line


def remove(slide, name: str) -> None:
    ph = _placeholder(slide, name)
    if ph is not None:
        ph._element.getparent().remove(ph._element)


def _photo(ctx: Ctx, slide, name: str, photo: dict | None) -> None:
    brief = place_photo(slide, PH_IDX[name], photo, ctx.base_dir, ctx.report.warnings)
    if brief:
        ctx.report.briefs.append((ctx.report.slides, brief))


def _box_px(ph) -> tuple[float, float, float, float]:
    return tuple(v / EMU_PER_PX for v in (ph.left, ph.top, ph.width, ph.height))


def _set_box(shape, x=None, y=None, w=None, h=None) -> None:
    """Platzhalter verschieben/strecken; alle vier Werte setzen, sonst fehlen geerbte Maße im xfrm."""
    left, top, width, height = shape.left, shape.top, shape.width, shape.height
    shape.left = left if x is None else px(x)
    shape.top = top if y is None else px(y)
    shape.width = width if w is None else px(w)
    shape.height = height if h is None else px(h)


def _quoted(text: str) -> str:
    text = text.strip()
    return text if text.startswith(("„", '"', "“", "»")) else f"„{text}“"


def _frame(ctx: Ctx, slide, s: dict, chapter: int | None = None) -> None:
    key = s["type"]
    dark = key in DARK_KEYS
    current = chapter if chapter is not None else s.get("chapter")
    if key not in NO_HEADER:
        header_nav(slide, ctx.chapters, current, dark, ctx.m)
    if (s.get("source") or s.get("legal")) and key not in ("photo", "cover"):
        source_row(slide, s.get("source"), s.get("legal"), dark, ctx.m)
    if key not in NO_FOOTER:
        footer(slide, ctx.pres["title"], ctx.report.slides, dark)
    if s.get("notes"):
        slide.notes_slide.notes_text_frame.text = s["notes"]


def _start(ctx: Ctx, s: dict, title: bool = True, default_title: str | None = None):
    slide = new_slide(ctx, LAYOUT_BY_KEY[s["type"]])
    fill(slide, "Kicker", s.get("kicker") or default_kicker(ctx.plan, s))
    if title:
        fill(slide, "Title", s.get("title") or default_title)
    return slide


def _points(slide, points: list[dict] | None) -> None:
    ph = _placeholder(slide, "Points")
    if ph is None:
        return
    if points:
        points_text(ph, points)
    else:
        ph._element.getparent().remove(ph._element)


def _focus(ctx: Ctx, slide, s: dict) -> None:
    focus = s["focus"]
    focus_box(slide, (1180, 340, 620), focus.get("label", "Im Fokus"), focus["title"], focus.get("text", ""), ctx.m)


# ---------------------------------------------------------------- Renderer

def render_cover(ctx: Ctx, s: dict) -> None:
    pres = ctx.pres
    presenter = pres.get("presenter") or {}
    slide = new_slide(ctx, LAYOUT_BY_KEY["cover"])
    fill(slide, "Kicker", s.get("kicker") or pres.get("kicker") or "ECOWERTH · Präsentation")
    fill(slide, "Title", s.get("title") or pres["title"])
    fill(slide, "Subtitle", pres.get("subtitle"))
    fill(slide, "Presenter", presenter.get("name"))
    fill(slide, "Credentials", "\n".join(t for t in (presenter.get("role"), pres.get("date")) if t))
    _frame(ctx, slide, s)


def render_agenda(ctx: Ctx, s: dict) -> None:
    slide = new_slide(ctx, LAYOUT_BY_KEY["agenda"])
    fill(slide, "Kicker", s.get("kicker") or "Agenda")
    fill(slide, "Title", s.get("title") or "Inhalt")
    items = s.get("items") or [{"title": name} for name in ctx.chapters]
    agenda_rows(slide, items, ctx.m)
    _frame(ctx, slide, s)


def _grow_chapter_title(ctx: Ctx, slide, text: str) -> None:
    """Lange Kapitelnamen brechen um: der Titel wächst nach unten, der Satz rückt nach."""
    title = _placeholder(slide, "Title")
    level = next(p for p in LAYOUT_BY_NAME[LAYOUT_BY_KEY["chapter"]].placeholders if p.name == "Title").levels[0]
    font = FONTS[level["font"]]
    para = [Para(text, level["size_pt"], level.get("ln_spc_pct", 100), font=font)]
    x, y, w, h = _box_px(title)
    need = ctx.m.block_height_px(para, w)
    if need <= h:
        return
    cap = 2 * ctx.m.line_height_px(level["size_pt"], level.get("ln_spc_pct", 100), font)
    if need > cap + 1 or ctx.m.long_words(text, level["size_pt"], w, font=font):
        ctx.report.warnings.append(f"Folie {ctx.report.slides}: Kapitelname „{text}“ passt nicht in zwei Zeilen "
                                   f"– kürzen oder trennen")
    extra = min(need, cap) - h
    _set_box(title, h=h + extra)
    line = _placeholder(slide, "Line")
    if line is not None:
        _set_box(line, y=_box_px(line)[1] + extra)


def render_chapter(ctx: Ctx, s: dict) -> None:
    chapters = ctx.chapters
    n = s["chapter"]
    slide = new_slide(ctx, LAYOUT_BY_KEY["chapter"])
    cx, cy, r, width = CHAPTER_RING
    motifs.progress_ring(slide.shapes, cx, cy, r, width, len(chapters), n)
    fill(slide, "Number", f"{n:02d}")
    fill(slide, "Kicker", s.get("kicker") or f"Kapitel {n} von {len(chapters)}")
    title = s.get("title") or chapters[n - 1]
    fill(slide, "Title", title)
    _grow_chapter_title(ctx, slide, title)
    fill(slide, "Line", s.get("line"))
    _frame(ctx, slide, s, chapter=n)


def render_hook(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s, title=False)
    fill(slide, "Quote", _quoted(s["quote"]))
    fill(slide, "TruthLabel", s["truth"].get("label", "Ausgangslage"))
    fill(slide, "TruthText", s["truth"]["text"])
    fill(slide, "AnswerLabel", s["answer"].get("label", "Unsere Antwort"))
    fill(slide, "AnswerText", s["answer"]["text"])
    _frame(ctx, slide, s)


def render_closing(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    closing_steps(slide, s["steps"], ctx.m)
    nxt = s["next"]
    fill(slide, "NextLabel", nxt.get("label", "Kontakt"))
    fill(slide, "NextTitle", nxt["title"])
    fill(slide, "NextText", nxt.get("text"))
    _frame(ctx, slide, s)


def render_points(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _points(slide, s["points"])
    if s.get("focus"):
        _set_box(_placeholder(slide, "Points"), w=980)
        _focus(ctx, slide, s)
    _frame(ctx, slide, s)


def render_points_photo(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _points(slide, s["points"])
    _photo(ctx, slide, "Photo", s["photo"])
    fill(slide, "Caption", s.get("caption"))
    _frame(ctx, slide, s)


def render_points_chart(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _points(slide, s["points"])
    bars_card(slide, s["chart"])
    _frame(ctx, slide, s)


def render_statement(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s, title=False)
    fill(slide, "Big", s.get("big"))
    fill(slide, "Statement", s.get("statement"))
    fill(slide, "Text", s.get("text"))
    _frame(ctx, slide, s)


def render_facts(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    fill(slide, "Lead", s.get("lead"))
    facts_band(slide, s["facts"], ctx.m)
    _frame(ctx, slide, s)


def render_brands(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    brand_columns(slide, s["columns"], ctx.m)
    _frame(ctx, slide, s)


def render_checklist(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    items = s["items"]
    if s.get("focus"):
        bottom = checklist(slide, items, ctx.m, w=980)
        _focus(ctx, slide, s)
    else:
        bottom = checklist(slide, items, ctx.m, columns=2 if len(items) > 5 else 1)
    if bottom > 900:
        ctx.report.warnings.append(f"Folie {ctx.report.slides}: Checkliste reicht bis {bottom:.0f} px – kürzen")
    _frame(ctx, slide, s)


def render_steps(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    start = s.get("start", 1)
    if any("photo" in step for step in s["steps"]):
        for brief in photo_steps(slide, s["steps"], ctx.base_dir, ctx.report.warnings, start, s.get("caption")):
            ctx.report.briefs.append((ctx.report.slides, brief))
    else:
        steps_grid(slide, s["steps"], ctx.m, start=start)
    _frame(ctx, slide, s)


def render_compare(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    compare_columns(slide, s["left"], s["right"], s.get("style", "do_dont"))
    _frame(ctx, slide, s)


def render_matrix(ctx: Ctx, s: dict) -> None:
    from .draw import style, textbox
    from .tokens import CONTENT_W, MARGIN_X, SOURCE_Y
    slide = _start(ctx, s)
    bottom = matrix_table(slide, s["columns"], s["rows"], ctx.m, corner=s.get("corner", ""))
    if s.get("note"):
        y = bottom + 12
        textbox(slide.shapes, "matrix-note", (MARGIN_X, y, CONTENT_W, 976 - y), [(s["note"], style(12, "muted"))])
    _frame(ctx, slide, s)


def render_scale(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    caption_y = 470
    if "scale" in s:
        scale_chips(slide, s["scale"]["items"], s["scale"].get("highlight", []))
    elif "timeline" in s:
        timeline(slide, s["timeline"]["points"])
        caption_y = 510                                    # unter zweizeiligen Zeitstrahl-Texten
    else:
        caption_y = tag_cloud(slide, s["tags"]["items"], s["tags"].get("highlight", []), ctx.m) + 22
    if s.get("caption"):
        scale_caption(slide, s["caption"], y=caption_y)
    points = _placeholder(slide, "Points")
    if points is not None and s.get("points"):
        top = max(600, caption_y + (70 if s.get("caption") else 30))
        _set_box(points, y=top, h=900 - top)
    _points(slide, s.get("points"))
    _frame(ctx, slide, s)


def render_photo(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _photo(ctx, slide, "Photo", s["photo"])
    fill(slide, "Caption", s.get("caption"))
    _frame(ctx, slide, s)


def render_photo_pair(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    _photo(ctx, slide, "PhotoLeft", s["left"]["photo"])
    _photo(ctx, slide, "PhotoRight", s["right"]["photo"])
    fill(slide, "LabelLeft", s["left"]["label"])
    fill(slide, "LabelRight", s["right"]["label"])
    _frame(ctx, slide, s)


def render_photo_annotated(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    box = _box_px(_placeholder(slide, "Photo"))   # Maße aus der Vorlage, auch wenn das Layout verändert wurde
    _photo(ctx, slide, "Photo", s["photo"])
    markers(slide, box, s["markers"], ctx.m)
    _frame(ctx, slide, s)


def render_case(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    fill(slide, "Story", s["story"])
    labels = tuple(s.get("columns") or ("Befund", "Ursache"))
    case_findings(slide, s["findings"], labels)
    _frame(ctx, slide, s)


def render_quote(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s, title=False)
    fill(slide, "Quote", _quoted(s["quote"]))
    fill(slide, "Name", s["name"])
    fill(slide, "Role", s.get("role"))
    if s.get("photo"):
        _photo(ctx, slide, "Photo", s["photo"])
    else:   # ohne Porträt rückt der Text nach links
        remove(slide, "Photo")
        for name in ("Quote", "Name", "Role"):
            ph = _placeholder(slide, name)
            if ph is not None:
                _set_box(ph, x=120, w=1560)
    _frame(ctx, slide, s)


def render_quiz(ctx: Ctx, s: dict) -> None:
    for reveal in (False, True):
        slide = new_slide(ctx, LAYOUT_BY_KEY["quiz"])
        fill(slide, "Kicker", s.get("kicker") or ("Wissenscheck · Auflösung" if reveal else "Wissenscheck"))
        fill(slide, "Question", s["question"])
        quiz_options(slide, s["options"], s["answer"] if reveal else None, s.get("explanation") if reveal else None)
        notes = dict(s)
        if reveal:
            notes["notes"] = s.get("explanation") or s.get("notes")
        _frame(ctx, slide, notes)


def render_summary(ctx: Ctx, s: dict) -> None:
    slide = _start(ctx, s)
    summary_items(slide, s["items"], ctx.m)
    _frame(ctx, slide, s)


RENDERERS: dict[str, Callable[[Ctx, dict], None]] = {
    "cover": render_cover, "agenda": render_agenda, "chapter": render_chapter, "hook": render_hook,
    "points": render_points, "points_photo": render_points_photo, "points_chart": render_points_chart,
    "statement": render_statement, "facts": render_facts, "brands": render_brands, "checklist": render_checklist,
    "steps": render_steps, "compare": render_compare, "matrix": render_matrix, "scale": render_scale,
    "photo": render_photo, "photo_pair": render_photo_pair, "photo_annotated": render_photo_annotated,
    "case": render_case, "quote": render_quote, "quiz": render_quiz, "summary": render_summary,
    "closing": render_closing,
}
