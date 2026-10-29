"""Erzeugt die ECOWERTH-Campus-Vorlage: Theme, Master und 23 Layouts."""
import json
from pathlib import Path
from xml.sax.saxutils import escape

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.opc.constants import CONTENT_TYPE as CT
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml import parse_xml
from pptx.oxml.ns import qn
from pptx.oxml.xmlchemy import OxmlElement
from pptx.parts.slide import SlideLayoutPart
from pptx.shapes.shapetree import SlideShapes

from . import draw, motifs, oxml
from .layouts import LAYOUTS, DecorDef, LayoutDef, PhDef, frame_decor
from .tokens import COLOR, PH_IDX, SLIDE_EMU, THEME_MAJOR, THEME_MINOR, px

THEME_NAME = "ECOWERTH Campus"
THEME_COLORS = {
    "dk1": COLOR["navy"], "lt1": COLOR["white"], "dk2": COLOR["navy_deep"], "lt2": COLOR["bg"],
    "accent1": COLOR["green"], "accent2": COLOR["blue"], "accent3": COLOR["red"], "accent4": COLOR["muted"],
    "accent5": COLOR["steel"], "accent6": COLOR["grey"], "hlink": COLOR["blue_ink"], "folHlink": COLOR["muted"],
}
FIRST_LAYOUT_ID = 2147483649

_NS = ('xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
       'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
       'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"')

LAYOUT_XML = (
    f'<p:sldLayout {_NS} preserve="1" userDrawn="1"><p:cSld name="{{name}}"><p:spTree>'
    '<p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>'
    '<p:grpSpPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="0" cy="0"/><a:chOff x="0" y="0"/>'
    '<a:chExt cx="0" cy="0"/></a:xfrm></p:grpSpPr></p:spTree></p:cSld>'
    '<p:clrMapOvr><a:masterClrMapping/></p:clrMapOvr></p:sldLayout>'
)


def _rpr(size: int, color: str, font: str) -> str:
    return (f'<a:defRPr sz="{size}" kern="1200"><a:solidFill><a:srgbClr val="{color}"/></a:solidFill>'
            f'<a:latin typeface="+{font}-lt"/><a:ea typeface="+{font}-ea"/><a:cs typeface="+{font}-cs"/></a:defRPr>')


def _ppr(level: int, size: int, color: str, font: str) -> str:
    return (f'<a:lvl{level}pPr marL="0" indent="0" algn="l" defTabSz="914400" rtl="0" eaLnBrk="1" '
            f'latinLnBrk="0" hangingPunct="1"><a:lnSpc><a:spcPct val="100000"/></a:lnSpc>'
            f'<a:spcBef><a:spcPts val="0"/></a:spcBef><a:buNone/>{_rpr(size, color, font)}</a:lvl{level}pPr>')


TXSTYLES_XML = (
    f'<p:txStyles {_NS}>'
    f'<p:titleStyle>{_ppr(1, 4000, COLOR["navy"], "mj")}</p:titleStyle>'
    f'<p:bodyStyle>{_ppr(1, 2200, COLOR["navy"], "mn")}{_ppr(2, 1700, COLOR["muted"], "mn")}'
    f'{_ppr(3, 1700, COLOR["muted"], "mn")}</p:bodyStyle>'
    f'<p:otherStyle><a:defPPr><a:defRPr lang="de-DE"/></a:defPPr>{_ppr(1, 1700, COLOR["navy"], "mn")}</p:otherStyle>'
    '</p:txStyles>'
)


def build_template(out: Path, assets_dir: Path) -> Path:
    prs = Presentation()
    prs.slide_width, prs.slide_height = SLIDE_EMU
    master = prs.slide_master
    _apply_theme(master)
    _style_master(master)
    _remove_layouts(master)
    for i, ldef in enumerate(LAYOUTS):
        layout = _new_layout(prs, master, ldef.name, FIRST_LAYOUT_ID + i)
        _populate(layout, ldef, Path(assets_dir))
    _set_language(prs)
    out = Path(out)
    out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(out)
    return out


def _set_language(prs) -> None:
    """Deutsch als Sprache für neue Texte (Rechtschreibprüfung in PowerPoint)."""
    default = prs.part._element.find(qn("p:defaultTextStyle"))
    if default is None:
        return
    for rpr in default.iter(qn("a:defRPr")):
        rpr.set("lang", "de-DE")


def _apply_theme(master) -> None:
    part = master.part.part_related_by(RT.THEME)
    xml = etree.fromstring(part.blob)
    xml.set("name", THEME_NAME)
    scheme = xml.find(".//" + qn("a:clrScheme"))
    scheme.set("name", THEME_NAME)
    for key, value in THEME_COLORS.items():
        slot = scheme.find(qn(f"a:{key}"))
        for child in list(slot):
            slot.remove(child)
        etree.SubElement(slot, qn("a:srgbClr")).set("val", value)
    fonts = xml.find(".//" + qn("a:fontScheme"))
    fonts.set("name", THEME_NAME)
    fonts.find(qn("a:majorFont")).find(qn("a:latin")).set("typeface", THEME_MAJOR)
    fonts.find(qn("a:minorFont")).find(qn("a:latin")).set("typeface", THEME_MINOR)
    part._blob = etree.tostring(xml, xml_declaration=True, encoding="UTF-8", standalone=True)


def _style_master(master) -> None:
    master.background.fill.solid()
    master.background.fill.fore_color.rgb = RGBColor.from_string(COLOR["white"])
    old = master._element.find(qn("p:txStyles"))
    new = parse_xml(TXSTYLES_XML)
    old.addprevious(new)
    master._element.remove(old)


def _remove_layouts(master) -> None:
    lst = master._element.get_or_add_sldLayoutIdLst()
    for el in list(lst):
        rId = el.get(qn("r:id"))
        lst.remove(el)
        master.part.drop_rel(rId)


def _new_layout(prs, master, name: str, layout_id: int):
    package = prs.part.package
    partname = package.next_partname("/ppt/slideLayouts/slideLayout%d.xml")
    element = parse_xml(LAYOUT_XML.format(name=escape(name, {'"': "&quot;"})))
    part = SlideLayoutPart(partname, CT.PML_SLIDE_LAYOUT, package, element)
    part.relate_to(master.part, RT.SLIDE_MASTER)
    rId = master.part.relate_to(part, RT.SLIDE_LAYOUT)
    entry = OxmlElement("p:sldLayoutId")
    entry.set("id", str(layout_id))
    entry.set(qn("r:id"), rId)
    master._element.get_or_add_sldLayoutIdLst().append(entry)
    return part.slide_layout


def background_of(frame: str) -> str:
    return COLOR["white"] if frame == "light" else COLOR["navy"]


def _populate(layout, ldef: LayoutDef, assets_dir: Path) -> None:
    layout.background.fill.solid()
    layout.background.fill.fore_color.rgb = RGBColor.from_string(background_of(ldef.frame))
    shapes = SlideShapes(layout.shapes._spTree, layout)
    for d in frame_decor(ldef.frame) + ldef.decor:
        _add_decor(shapes, d, assets_dir)
    for ph in ldef.placeholders:
        _add_placeholder(shapes, ph)


def _add_decor(shapes, d: DecorDef, assets_dir: Path) -> None:
    st = d.style
    if d.kind == "image":
        draw.picture(shapes, d.name, assets_dir / st["file"], d.box)
    elif d.kind == "text":
        draw.textbox(shapes, d.name, d.box, [(st["text"], st["level"])], anchor=st.get("anchor", "t"))
    elif d.kind == "hline":
        x, y, w, h = d.box
        draw.hline(shapes, d.name, x, y, w, st["color"], h)
    elif d.kind == "kreislauf":
        cx, cy, r, _ = d.box
        motifs.kreislauf(shapes, cx, cy, r, prefix=d.name, core_fill=st.get("core_fill", "white"),
                         dot_color=st.get("dot_color", "white"), hair_color=st.get("hair_color", "grey"))
    else:
        kind = {"rect": MSO_SHAPE.RECTANGLE, "ellipse": MSO_SHAPE.OVAL}[d.kind]
        draw.shape(shapes, kind, d.name, d.box, fill=st.get("fill"), line=st.get("line"))


def _add_placeholder(shapes, ph: PhDef) -> None:
    x, y, w, h = (px(v) for v in ph.box)
    if ph.ph_type == "pic":
        shape = shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    else:
        shape = shapes.add_textbox(x, y, w, h)
    idx = None if ph.ph_type == "title" else PH_IDX[ph.name]
    oxml.to_placeholder(shape, ph.ph_type, idx, ph.name, ph.levels, ph.prompt,
                        ph.anchor, ph.insets, ph.geom)
    fill = ph.fill or (COLOR["surface2"] if ph.ph_type == "pic" else None)
    draw.apply_fill_line(shape, fill, ph.line, ph.fill_alpha)


def layout_spec() -> dict:
    """Maße und Stile aller Layouts als JSON-taugliches Dict (Quelle für Prüfung und Skill-Doku)."""
    spec = {}
    for ldef in LAYOUTS:
        spec[ldef.name] = {
            "key": ldef.key,
            "frame": ldef.frame,
            "placeholders": {
                ph.name: {
                    "idx": 0 if ph.ph_type == "title" else PH_IDX[ph.name],
                    "type": ph.ph_type,
                    "box": list(ph.box),
                    "levels": ph.levels,
                    "insets": list(ph.insets),
                }
                for ph in ldef.placeholders
            },
        }
    return json.loads(json.dumps(spec))


def write_layout_spec(path: Path) -> None:
    Path(path).write_text(json.dumps(layout_spec(), indent=2, ensure_ascii=False), encoding="utf-8")
